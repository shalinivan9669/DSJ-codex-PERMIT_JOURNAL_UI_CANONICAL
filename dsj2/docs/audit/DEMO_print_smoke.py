"""Offline audit harness; synthetic data only; no database/network access."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from io import BytesIO
import base64
import json
import os
import re
import subprocess
import sys
from PIL import Image
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / "DEMO_print_smoke"
OUT.mkdir(exist_ok=True)
ENV = {**os.environ, "PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8", "PYTHONDONTWRITEBYTECODE": "1"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
results = []

def run(name, script, template, payload):
    output = OUT / name
    args = [sys.executable, "-X", "utf8", str(ROOT / "scripts" / script)]
    if template:
        args.append(str(ROOT / "docs" / "experimental" / template))
    args.append(str(output))
    result = subprocess.run(args, input=json.dumps(payload, ensure_ascii=False).encode(), capture_output=True, env=ENV, timeout=30)
    detail = {"output": name, "exit_code": result.returncode, "stderr": result.stderr.decode("utf8", errors="replace")}
    if result.returncode == 0 and name.endswith(".docx"):
        with ZipFile(output) as archive:
            xml = archive.read("word/document.xml")
            root = ET.fromstring(xml)
            text = "\n".join(n.text or "" for n in root.iter(W + "t"))
            (OUT / (name + ".text.txt")).write_text(text, encoding="utf8")
            detail.update(bytes=output.stat().st_size, synthetic_name_present="Тестов" in text, kazakh_letters_present="Ә" in text, remaining_placeholders=re.findall(r"\{\{[A-Z0-9_]+\}\}", text), external_relationships=sum(archive.read(n).count(b'TargetMode="External"') for n in archive.namelist() if n.endswith(".rels")), media_files=len([n for n in archive.namelist() if n.startswith("word/media/")]), manual_page_breaks=len(root.findall('.//' + W + 'br[@' + W + 'type="page"]')), mail_merge_settings=(b"mailMerge" in archive.read("word/settings.xml") if "word/settings.xml" in archive.namelist() else False))
    results.append(detail)
    return output

fields = {"Берілді": "Тестов Тест Тестович", "В_том_что_он": "Синтетическое обучение DEMO", "ГОД": "2026", "День_месяц": "22.09", "Должность": "Тестовый инженер", "Жұмыс__орны_": "DEMO сынақ ұйымы", "Лауазымы": "Әзірлеуші Өңдеуші", "Место_работы__": "DEMO тестовая организация", "Номер_серии": "DEMO", "Номер_удостоверения": "DEMO-БТ-00001", "Протокол_": "DEMO-ПТ-00001", "ФИО": "Тестов Тест Тестович"}
run("DEMO_biot_single.docx", "generate_biot_card.py", "biot/biot-card-template.docx", {"fields": fields})
second = {**fields, "ФИО": "Тестов Оченьдлинноеимядляпроверкипереноса Оченьдлинноеотчестводляпроверкипереноса", "Берілді": "Әділбек Өмірсерік Қанатұлы", "Номер_удостоверения": "DEMO-БТ-00002", "Протокол_": "DEMO-ПТ-00002"}
run("DEMO_biot_batch_2.docx", "generate_biot_mail_merge_bundle.py", "biot/biot-card-template.docx", {"rows": [{"fields": fields}, {"fields": second}]})

photo_buffer = BytesIO()
Image.new("RGB", (32, 40), (10, 100, 150)).save(photo_buffer, format="PNG")
photo = {"dataUrl": "data:image/png;base64," + base64.b64encode(photo_buffer.getvalue()).decode(), "slot": {"mode": "floating_rect", "shapeId": "DSJPhotoSlotPB", "style": "position:absolute;margin-left:-2.1pt;margin-top:121.5pt;width:59.8pt;height:79.65pt;z-index:251709952;visibility:visible;mso-wrap-style:square;mso-position-horizontal:absolute;mso-position-horizontal-relative:text;mso-position-vertical:absolute;mso-position-vertical-relative:page"}}
pb_fields = {'"Месяц"': "09", "Год": "2026", "Действительно_Год": "2027", "День": "22", "Жұмыс_орны_лауазымы": "DEMO тестовая организация/DEMO сынақ ұйымы", "Месяц": "09", "Номер_удостоверения": "DEMO-ПБ-00001", "Прослушала_курс": "Синтетическое обучение DEMO", "Протокол_": "DEMO-ПБ-ПТ-00001", "ФИО": "Тестов Тест Тестович", "должность": "Тестовый инженер/Әзірлеуші Өңдеуші"}
run("DEMO_pb_photo.docx", "generate_biot_card.py", "pb/pb-card-template.docx", {"fields": pb_fields, "photo": photo})

template_path = ROOT / "docs/experimental/ps/ps-witness-certificate-template.docx"
with ZipFile(template_path) as archive:
    template_text = "".join(n.text or "" for n in ET.fromstring(archive.read("word/document.xml")).iter(W + "t"))
witness_fields = {key: "DEMO" for key in set(re.findall(r"\{\{[A-Z0-9_]+\}\}", template_text))}
witness_fields.update({"{{KB_NUMBER}}": "КБ № DEMO-00001", "{{REGISTRATION_NUMBER}}": "DEMO-00001", "{{FULL_NAME_RU}}": "Тестов Тест Тестович", "{{FULL_NAME_KZ}}": "Әділбек Өмірсерік Қанатұлы", "{{PROFESSION_RU}}": "Тестовый инженер", "{{PROFESSION_KZ}}": "Әзірлеуші Өңдеуші", "{{PROTOCOL_NUMBER_DISPLAY}}": "DEMO-ПТ-00001", "{{EDU_ORG_RU}}": "DEMO тестовый центр", "{{EDU_ORG_KZ}}": "DEMO сынақ орталығы"})
run("DEMO_ps_witness.docx", "generate_ps_witness_certificate.py", "ps/ps-witness-certificate-template.docx", {"rows": [{"fields": witness_fields}]})

registry = run("DEMO_registry_formula_probe.xlsx", "export_card_request_registry.py", None, {"request": {"title": "DEMO synthetic audit", "certificateTypeLabel": "БиОТ"}, "items": [{"fullName": "=1+1", "certificateNumber": "DEMO-00001", "protocolNumber": "DEMO-ПТ-00001"}]})
if registry.exists():
    cell = load_workbook(registry, data_only=False)["Реестр"]["D2"]
    results[-1].update(probe_cell="Реестр!D2", probe_value=cell.value, probe_data_type=cell.data_type, formula_injection_confirmed=(cell.data_type == "f"))

inventory = []
for path in (ROOT / "docs/experimental").glob("**/*template.docx"):
    if "correspondence" in path.as_posix():
        continue
    with ZipFile(path) as archive:
        tree = ET.fromstring(archive.read("word/document.xml"))
        text = "\n".join(n.text or "" for n in tree.iter(W + "t"))
        sample_paragraphs = ["".join(n.text or "" for n in p.iter(W + "t")) for p in tree.iter(W + "p")]
        inventory.append({"template": str(path.relative_to(ROOT)), "bytes": path.stat().st_size, "old_center_reference_count": text.count("Стандарт"), "has_otcenter_literal": bool(re.search(r"otcenter|ot center", text, re.I)), "provider_or_signer_paragraphs": [p for p in sample_paragraphs if any(s in p for s in ["Стандарт", "Баянов", "Жакибеков", "Флеглер", "Есен"])], "merge_instructions": sorted(set(n.text or "" for n in tree.iter(W + "instrText") if "MERGEFIELD" in (n.text or "")))})
(OUT / "DEMO_template_inventory.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf8")
(OUT / "DEMO_smoke_results.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf8")
print(json.dumps(results, ensure_ascii=False, indent=2))
