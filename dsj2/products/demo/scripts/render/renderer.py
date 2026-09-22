"""DEMO render boundary. stdin JSON -> one output plus bounded JSON metadata.
All paths are supplied by trusted API/worker, never by an uploaded document.
"""
from __future__ import annotations
import base64
import csv
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import importlib.metadata
import threading
from copy import deepcopy
from datetime import date
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo
from lxml import etree as E
from PIL import Image, ImageOps, ImageFont
from openpyxl import Workbook, load_workbook
from sanitize_templates import replace_text_nodes, deterministic_zip
from package_xml import normalize_package
from xml_input import validate_xlsx_xml

ROOT=Path(__file__).resolve().parents[2]
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
PKG='{http://schemas.openxmlformats.org/package/2006/relationships}'
NS={'w':W[1:-1]}
RENDERER_VERSION='demo-ooxml-6/libreoffice-26.2.6.3'
RU=['января','февраля','марта','апреля','мая','июня','июля','августа','сентября','октября','ноября','декабря']
KZ=['қаңтар','ақпан','наурыз','сәуір','мамыр','маусым','шілде','тамыз','қыркүйек','қазан','қараша','желтоқсан']
Image.MAX_IMAGE_PIXELS=20_000_000
_converter_versions={}
_converter_version_lock=threading.Lock()

def date_parts(value):
    if not value: return dict.fromkeys(['DAY','MONTH','YEAR','YEAR_SHORT','DAY_MONTH','DATE','MONTH_RU','MONTH_KZ'],'')
    d=date.fromisoformat(value)
    return {'DAY':f'{d.day:02}','MONTH':f'{d.month:02}','YEAR':str(d.year),'YEAR_SHORT':str(d.year)[2:],
            'DAY_MONTH':f'{d.day:02}.{d.month:02}','DATE':f'{d.day:02}.{d.month:02}.{d.year}',
            'MONTH_RU':RU[d.month-1],'MONTH_KZ':KZ[d.month-1]}

def fields_for(snapshot,item):
    assignment=item['assignment']; issuer=snapshot['issuer']
    commission=issuer.get('commission',[])
    member=lambda i: (' — '.join(filter(None,[commission[i].get('name'),commission[i].get('position')])) if len(commission)>i else '')
    both=lambda a,b: ' / '.join(filter(None,[item.get(a),item.get(b)]))
    numbers=item.get('numbers',{})
    number=item.get('number','')
    credential_number=item.get('credentialNumber','') if snapshot['templateId']=='ps-protocol' else number
    fields={'EMPTY':'','ONE':'1','TWO':'2','SERIES':'','FULL_NAME_RU':item.get('fullNameRu',''),
        'FULL_NAME_KZ':item.get('fullNameKz',''),'FULL_NAME_BOTH':both('fullNameRu','fullNameKz'),
        'POSITION_RU':item.get('positionRu',''),'POSITION_KZ':item.get('positionKz',''),
        'POSITION_BOTH':both('positionRu','positionKz'),'WORKPLACE_RU':item.get('workplaceRu',''),
        'WORKPLACE_KZ':item.get('workplaceKz',''),'WORKPLACE_BOTH':both('workplaceRu','workplaceKz'),
        'NUMBER':credential_number,'KB_NUMBER':number,'REGISTRATION_NUMBER':item.get('registrationNumber',''),
        'PROTOCOL_NUMBER':item.get('protocolNumber') or numbers.get('protocol') or (number if 'protocol' in snapshot['templateId'] else assignment.get('externalBasisNumber','')),
        'SUBJECT':assignment.get('trainingSubject',''),'RESULT':assignment.get('result',''),
        'REASON':assignment.get('reason',''),'EDUCATION':assignment.get('education',''),'HOURS':str(assignment.get('hours','')),
        'ISSUER_RU':issuer.get('nameRu',''),'ISSUER_KZ':issuer.get('nameKz',''),
        'ISSUER_BOTH':' / '.join(filter(None,[issuer.get('nameRu'),issuer.get('nameKz')])),
        'EDU_ORG_RU':issuer.get('nameRu',''),'EDU_ORG_KZ':issuer.get('nameKz',''),
        'CITY_RU':issuer.get('cityRu',''),'CITY_KZ':issuer.get('cityKz',''),
        'ADDRESS_RU':issuer.get('addressRu',''),'ADDRESS_KZ':issuer.get('addressKz',''),
        'CHAIR':member(0),'MEMBER_1':member(1),'MEMBER_2':member(2),'APPROVAL_BASIS':issuer.get('approvalBasis','')}
    for prefix,key in [('DOCUMENT','documentDate'),('PROTOCOL','protocolDate'),('VALID','validUntil'),('TRAINING_START','trainingStart'),('TRAINING_END','trainingEnd'),('ISSUE','documentDate')]:
        for part,value in date_parts(assignment.get(key)).items(): fields[f'{prefix}_{part}']=value
    fields.update(PROFESSION_RU=fields['POSITION_RU'],PROFESSION_KZ=fields['POSITION_KZ'],PROTOCOL_NUMBER_DISPLAY=fields['PROTOCOL_NUMBER'])
    from biot_2026 import current_fields
    fields.update(current_fields(snapshot,item))
    fields['TRAINING_START_YEAR_FULL']=fields['TRAINING_START_YEAR']; fields['TRAINING_END_YEAR_FULL']=fields['TRAINING_END_YEAR']
    return fields

def photo_normalize(input_path,output_path,options):
    if Path(input_path).stat().st_size>5*1024*1024: raise ValueError('PHOTO_BYTES_LIMIT')
    with Image.open(input_path) as check:
        if check.format not in ['PNG','JPEG']: raise ValueError('PHOTO_FORMAT')
        if check.width*check.height>20_000_000: raise ValueError('PHOTO_PIXEL_LIMIT')
        check.verify()
    with Image.open(input_path) as original:
        image=ImageOps.exif_transpose(original).convert('RGB')
        rotation=int(options.get('rotation',0))
        if rotation not in [0,90,180,270]: raise ValueError('PHOTO_ROTATION')
        if rotation: image=image.rotate(-rotation,expand=True)
        crop=options.get('crop')
        if crop:
            x,y,w,h=[float(crop[k]) for k in ['x','y','width','height']]
            if min(x,y)<0 or min(w,h)<=0 or x+w>1 or y+h>1: raise ValueError('PHOTO_CROP')
            image=image.crop((round(x*image.width),round(y*image.height),round((x+w)*image.width),round((y+h)*image.height)))
        if min(image.size)<96: raise ValueError('PHOTO_TOO_SMALL')
        warning='LOW_PRINT_RESOLUTION' if image.width<360 or image.height<480 else None
        image.thumbnail((1200,1600))
        image.save(output_path,'PNG',optimize=True)
        return {'width':image.width,'height':image.height,'mimeType':'image/png','warning':warning}

def add_mark(files,text):
    root=E.fromstring(files['word/document.xml'])
    header=E.Element(W+'hdr',nsmap={'w':W[1:-1]})
    p=E.SubElement(header,W+'p'); pr=E.SubElement(p,W+'pPr'); E.SubElement(pr,W+'jc',{W+'val':'center'})
    run=E.SubElement(p,W+'r'); props=E.SubElement(run,W+'rPr'); E.SubElement(props,W+'sz',{W+'val':'14'}); E.SubElement(props,W+'color',{W+'val':'9C2020'})
    E.SubElement(run,W+'t').text=text
    files['word/demo-header.xml']=E.tostring(header,xml_declaration=True,encoding='utf-8')
    rels=E.fromstring(files['word/_rels/document.xml.rels'])
    rid='rIdDemoHeader'; E.SubElement(rels,PKG+'Relationship',Id=rid,Type=R[1:-1]+'/header',Target='demo-header.xml')
    for section in root.iter(W+'sectPr'):
        for old in section.findall(W+'headerReference'): section.remove(old)
        for kind in ['default','first','even']:
            section.insert(0,E.Element(W+'headerReference',{W+'type':kind,R+'id':rid}))
    files['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='utf-8')
    ct=E.fromstring(files['[Content_Types].xml']); E.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',PartName='/word/demo-header.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml')
    files['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='utf-8')
    files['word/document.xml']=E.tostring(root,xml_declaration=True,encoding='utf-8')

def fit_textboxes(tree,tid,has_photo):
    """Keep the historic panel/section geometry and reserve the original photo slot.
    Source boxes used spaces/tabs as positioning, which clipped substituted text.
    Reflow populated boxes only; blank repeat-examination panels stay unchanged.
    """
    if tid not in ['biot-worker-card','ptm-card','pb-card','ps-card']:return
    for box in tree.iter(W+'txbxContent'):
        text=''.join(box.itertext())
        if '{{' not in text: continue
        shape=next((n for n in box.iterancestors() if E.QName(n).localname in ['shape','rect','anchor']),None)
        issuer_strip='{{ISSUER_BOTH}}' in text
        style=shape.get('style','') if shape is not None else ''
        width_match=re.search(r'(?:^|;)width:([\d.]+)pt',style)
        height_match=re.search(r'(?:^|;)height:([\d.]+)pt',style)
        width=float(width_match[1]) if width_match else 260
        height=float(height_match[1]) if height_match else 160
        if shape is not None and E.QName(shape).localname=='anchor':
            extent=next((n for n in shape if E.QName(n).localname=='extent'),None)
            if extent is not None:width=int(extent.get('cx','3302000'))/12700;height=int(extent.get('cy','2032000'))/12700
        # Data-bearing first panel has the photo slot, other panels use their full width.
        photo_panel=has_photo and 'FULL_NAME' in text
        if tid=='ps-card' and 'FULL_NAME' in text and shape is not None:
            if E.QName(shape).localname=='shape':
                shape.set('style',re.sub(r'margin-top:([-.\d]+)pt',lambda m:'margin-top:'+str(float(m[1])+8)+'pt',style))
            else:
                for position in shape.iter():
                    if E.QName(position).localname=='positionV':
                        offset=next((n for n in position if E.QName(n).localname=='posOffset'),None)
                        if offset is not None:offset.text=str(int(offset.text or '0')+8*12700)
        paragraphs=[]
        for p in list(box):
            if p.tag!=W+'p': continue
            nodes=list(p.iter(W+'t'));value=re.sub(r'\s+',' ',''.join(n.text or '' for n in nodes)).strip()
            if value:paragraphs.append((p,value))
            elif not any(E.QName(n).localname in ['drawing','pict','br'] for n in p.iter()):box.remove(p)
        for p,value in paragraphs:
            # Right/left insets apply only to the fields below the panel title.
            is_title=('NUMBER' in value and any(title in value.upper() for title in ['УДОСТОВЕРЕНИЕ','КУӘЛІК','КУƏЛІК']))
            inset=24 if tid=='biot-worker-card' and 'FULL_NAME' in text else 18 if tid=='ptm-card' and 'FULL_NAME' in text else 100 if tid=='pb-card' and 'FULL_NAME' in text and not is_title else 88 if photo_panel and tid=='ps-card' and not is_title else 0
            if tid=='ps-card' and 'Тапсырылған емтихандар' in text:inset=24
            right=76 if photo_panel and tid=='ptm-card' and not is_title and 'ISSUER' not in value else 0
            for child in list(p):p.remove(child)
            pr=E.SubElement(p,W+'pPr');E.SubElement(pr,W+'spacing',{W+'before':'0',W+'after':'0',W+'line':'180',W+'lineRule':'exact'})
            E.SubElement(pr,W+'ind',{W+'left':str(round(inset*20)),W+'right':str(round(right*20)),W+'firstLine':'0'})
            E.SubElement(pr,W+'jc',{W+'val':'center' if issuer_strip or (is_title and tid!='ptm-card') else 'left'})
            run=E.SubElement(p,W+'r');rp=E.SubElement(run,W+'rPr')
            E.SubElement(rp,W+'rFonts',{W+'ascii':'Liberation Serif',W+'hAnsi':'Liberation Serif',W+'cs':'Liberation Serif'})
            E.SubElement(rp,W+'sz',{W+'val':'16'});E.SubElement(rp,W+'szCs',{W+'val':'16'})
            if is_title:E.SubElement(rp,W+'b')
            E.SubElement(run,W+'t').text=value
        if tid=='ps-card' and 'FULL_NAME' in text:
            spacer=E.Element(W+'p');pr=E.SubElement(spacer,W+'pPr');E.SubElement(pr,W+'spacing',{W+'line':'160',W+'lineRule':'exact'});box.insert(0,spacer)
        box.set('data-demo-width',str(width));box.set('data-demo-height',str(height))
        if issuer_strip:box.set('data-demo-bottom-pad','2')
        if photo_panel:
            box.set('data-demo-photo-band',{'ptm-card':'48,124','pb-card':'25,116','ps-card':'42,138'}.get(tid,'0,0'))
            box.set('data-demo-photo-side','right' if tid=='ptm-card' else 'left')

def fit_rendered_textboxes(tree):
    for box in tree.iter(W+'txbxContent'):
        if 'data-demo-width' not in box.attrib:continue
        width=float(box.attrib.pop('data-demo-width'));height=float(box.attrib.pop('data-demo-height'))
        bottom_pad=float(box.attrib.pop('data-demo-bottom-pad','6'))
        band=box.attrib.pop('data-demo-photo-band',None);side=box.attrib.pop('data-demo-photo-side',None)
        for size in [8]:
            font=ImageFont.truetype(str(ROOT/'assets/fonts/LiberationSerif-Regular.ttf'),round(size*10))
            total=0;unbreakable_overflow=False
            for p in box.findall(W+'p'):
                if band:
                    start,end=map(float,band.split(','));ind=p.find(W+'pPr/'+W+'ind')
                    # The old uniform inset wasted the full width below the photo.
                    # Release it only after the paragraph has cleared the image.
                    if ind is not None and total>=end:
                        ind.set(W+side,'360' if side=='right' else '600')
                ind=p.find(W+'pPr/'+W+'ind');avail=width-9-((int(ind.get(W+'left','0'))+int(ind.get(W+'right','0')))/20 if ind is not None else 0)
                lines=1;current=''
                for word in ''.join(n.text or '' for n in p.iter(W+'t')).split():
                    if font.getlength(word)/10>avail:unbreakable_overflow=True
                    candidate=(current+' '+word).strip()
                    if current and font.getlength(candidate)/10>avail:lines+=1;current=word
                    else:current=candidate
                total+=lines*(size+1)
            if total<=height-bottom_pad and not unbreakable_overflow:break
        if total>height-bottom_pad or unbreakable_overflow:raise ValueError('PRINT_LAYOUT_OVERFLOW')
        for node in box.iter():
            if node.tag in [W+'sz',W+'szCs']:node.set(W+'val',str(round(size*2)))
            if node.tag==W+'spacing':node.set(W+'line',str(round((size+1)*20)))

def fit_certificate_name(tree,value):
    # The historic certificate's name box extends behind the subject panels.
    # Keep its geometry and limit the name to the clear 60pt band above them.
    for paragraph in tree.iter(W+'p'):
        if '{{FULL_NAME_RU}}' not in ''.join(paragraph.itertext()):continue
        for size in [33,32,30,28,26,24,22,20,18,16,14,12]:
            font=ImageFont.truetype(str(ROOT/'assets/fonts/LiberationSerif-Regular.ttf'),size*10)
            lines=1;current='';overflow=False
            for word in value.split():
                if font.getlength(word)/10>530:overflow=True
                candidate=(current+' '+word).strip()
                if current and font.getlength(candidate)/10>530:lines+=1;current=word
                else:current=candidate
            if not overflow and lines*(size+1)<=60:break
        if overflow or lines*(size+1)>60:raise ValueError('PRINT_LAYOUT_OVERFLOW')
        pr=paragraph.find(W+'pPr')
        if pr is None:pr=E.Element(W+'pPr');paragraph.insert(0,pr)
        for old in pr.findall(W+'spacing'):pr.remove(old)
        E.SubElement(pr,W+'spacing',{W+'before':'0',W+'after':'0',W+'line':str((size+1)*20),W+'lineRule':'exact'})
        for run in paragraph.findall(W+'r'):
            rp=run.find(W+'rPr')
            if rp is None:rp=E.Element(W+'rPr');run.insert(0,rp)
            for key in ['sz','szCs']:
                for old in rp.findall(W+key):rp.remove(old)
                E.SubElement(rp,W+key,{W+'val':str(size*2)})

def render_one(snapshot,item,template):
    with ZipFile(template) as z: files={n:z.read(n) for n in z.namelist()}
    fields=fields_for(snapshot,item)
    # Shared with contracts PRINT_LIMITS.maxUnbroken. Reject before registration via
    # preflight; never truncate a legal name or solve overflow with a micro-font.
    if any(len(token)>80 for value in fields.values() for token in str(value).split()):
        raise ValueError('PRINT_LAYOUT_OVERFLOW')
    # Some historic forms have one bilingual cell backed by a RU-only merge key.
    # Decide from the immutable template bytes, never from the current manifest.
    available=set(re.findall(r'\{\{([A-Z0-9_]+)\}\}', ''.join(''.join(E.fromstring(data).itertext()) for name,data in files.items() if name.endswith('.xml'))))
    for base in ['FULL_NAME','POSITION','WORKPLACE']:
        if base+'_RU' in available and not ({base+'_KZ',base+'_BOTH'} & available):
            fields[base+'_RU']=fields[base+'_BOTH']
    for name,data in list(files.items()):
        if not name.endswith('.xml'): continue
        root=E.fromstring(data)
        for node in root.iter():
            for attribute in list(node.attrib):
                if E.QName(attribute).localname=='gfxdata':del node.attrib[attribute]
        for p in root.iter(W+'p'):
            nodes=p.xpath('./w:r/w:t | ./w:hyperlink/w:r/w:t',namespaces=NS)
            label=''.join(n.text or '' for n in nodes)
            if label.count('{{ISSUER_RU}}')>1:
                nodes[0].text='{{ISSUER_RU}}'
                for n in nodes[1:]:n.text=''
            else:replace_text_nodes(nodes,r'ТОО\s+Аттестац\w*','')
        if snapshot['templateId']=='ps-card':
            for p in root.iter(W+'p'):
                nodes=p.xpath('./w:r/w:t',namespaces=NS);label=' '.join(''.join(n.text or '' for n in nodes).split())
                if '{{RESULT}}' in label and ('Выпускной' in label or 'Практическое' in label):
                    nodes[0].text='Результат / Нәтиже: {{RESULT}}' if 'Выпускной' in label else ''
                    for n in nodes[1:]:n.text=''
            for table in root.iter(W+'tbl'):
                if '{{ONE}}' not in ''.join(table.itertext()):continue
                position=table.find(W+'tblPr/'+W+'tblpPr')
                if position is not None:position.set(W+'vertAnchor','page');position.set(W+'tblpY','3000')
                for row in table.findall(W+'tr'):
                    if '{{TWO}}' in ''.join(row.itertext()):
                        for n in row.iter(W+'t'):n.text=''
                for p in table.iter(W+'p'):
                    nodes=list(p.iter(W+'t'))
                    for token in ['SUBJECT','RESULT']:replace_text_nodes(nodes,re.escape('{{'+token+'}} / {{'+token+'}}'),'{{'+token+'}}')
                    pr=p.find(W+'pPr')
                    if pr is None:pr=E.Element(W+'pPr');p.insert(0,pr)
                    for old in pr.findall(W+'spacing'):pr.remove(old)
                    E.SubElement(pr,W+'spacing',{W+'before':'0',W+'after':'0',W+'line':'180',W+'lineRule':'exact'})
                    for run in p.findall(W+'r'):
                        rp=run.find(W+'rPr')
                        if rp is None:rp=E.Element(W+'rPr');run.insert(0,rp)
                        for old in rp.findall(W+'sz'):rp.remove(old)
                        E.SubElement(rp,W+'sz',{W+'val':'16'})
        fit_textboxes(root,snapshot['templateId'],bool(item.get('photoAssetId')))
        current_biot=any((n.get(W+'val') or '').startswith('BIOT2026_') for n in root.iter(W+'tblCaption'))
        if snapshot['templateId']=='biot-itr-certificate' and not current_biot:fit_certificate_name(root,fields['FULL_NAME_RU'])
        if snapshot['templateId']=='pb-protocol':
            body=root.find(W+'body');previous_blank=False
            for paragraph in list(body) if body is not None else []:
                blank=paragraph.tag==W+'p' and not ''.join(paragraph.itertext()).strip() and not any(E.QName(n).localname in ['drawing','pict','sectPr','br'] for n in paragraph.iter())
                if blank and previous_blank:body.remove(paragraph)
                previous_blank=blank
        for p in root.iter(W+'p'):
            nodes=p.xpath('./w:r/w:t | ./w:hyperlink/w:r/w:t',namespaces=NS)
            for key,value in fields.items(): replace_text_nodes(nodes,re.escape('{{'+key+'}}'),str(value))
        leftovers=re.findall(r'\{\{[A-Z0-9_]+\}\}', ''.join(root.itertext()))
        if leftovers: raise ValueError('TEMPLATE_FIELDS_MISSING:'+','.join(sorted(set(leftovers))))
        fit_rendered_textboxes(root)
        if current_biot:
            from biot_2026 import assert_page_fit
            assert_page_fit(root)
        files[name]=E.tostring(root,xml_declaration=True,encoding='utf-8')
    photo_id=item.get('photoAssetId'); tid=snapshot['templateId']
    if photo_id and tid in ['ptm-card','pb-card','ps-card']:
        key=snapshot.get('photos',{}).get(photo_id)
        if not key: raise ValueError('PHOTO_ASSET_MISSING')
        store=Path(os.environ['DEMO_ARTIFACT_ROOT']).resolve(); photo=(store/key).resolve()
        if not photo.is_relative_to(store): raise ValueError('PHOTO_PATH')
        coords={'ptm-card':(215.4,123.75,51,67.5),'pb-card':(-2.1,121.5,59.8,79.65),'ps-card':(8,108,56,74)}[tid]
        x,y,w,h=coords
        blob=photo.read_bytes();photo_digest=hashlib.sha256(blob).hexdigest()
        expected=re.search(r'/(\w{64})-[0-9a-f-]{36}\.png$',key)
        if expected and expected[1]!=photo_digest:raise ValueError('PHOTO_HASH_MISMATCH')
        media='photo-'+photo_digest+'.png';rid='rIdDemoPhoto'
        tree=E.fromstring(files['word/document.xml']);first=tree.find(W+'body/'+W+'p')
        run=E.SubElement(first,W+'r');pict=E.SubElement(run,W+'pict')
        rect=E.SubElement(pict,'{urn:schemas-microsoft-com:vml}rect',id='DemoPhoto',stroked='f',filled='t',style=f'position:absolute;margin-left:{x}pt;margin-top:{y}pt;width:{w}pt;height:{h}pt;z-index:251709952;visibility:visible;mso-position-horizontal:absolute;mso-position-horizontal-relative:text;mso-position-vertical:absolute;mso-position-vertical-relative:page')
        E.SubElement(rect,'{urn:schemas-microsoft-com:vml}imagedata',{R+'id':rid})
        rels=E.fromstring(files['word/_rels/document.xml.rels']);E.SubElement(rels,PKG+'Relationship',Id=rid,Type=R[1:-1]+'/image',Target='media/'+media)
        ct=E.fromstring(files['[Content_Types].xml'])
        if not any(n.get('Extension')=='png' for n in ct):E.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='png',ContentType='image/png')
        files.update({'word/document.xml':E.tostring(tree,xml_declaration=True,encoding='utf-8'),'word/_rels/document.xml.rels':E.tostring(rels,xml_declaration=True,encoding='utf-8'),'[Content_Types].xml':E.tostring(ct,xml_declaration=True,encoding='utf-8'),'word/media/'+media:blob})
    if snapshot.get('demoMode'): add_mark(files,'ДЕМО — НЕ ЯВЛЯЕТСЯ ВЫДАННЫМ ДОКУМЕНТОМ')
    elif snapshot.get('mode')=='draft-preview': add_mark(files,'ПРЕДПРОСМОТР — НЕ ЯВЛЯЕТСЯ ВЫДАННЫМ ДОКУМЕНТОМ')
    tree=E.fromstring(files['word/document.xml']);body=tree.find(W+'body')
    while len(body)>1:
        tail=body[-2] if body[-1].tag==W+'sectPr' else body[-1]
        if tail.tag!=W+'p' or ''.join(tail.itertext()).strip() or tail.find('.//'+W+'sectPr') is not None or any(E.QName(n).localname in ['drawing','pict','br'] for n in tail.iter()):break
        body.remove(tail)
    # Historical BIOT used 200 empty paragraphs as a page break. Replace only such
    # runs with a real break so concatenated recipients cannot move anchored panels.
    streak=[]
    for node in list(body):
        blank=node.tag==W+'p' and not any((t.text or '').strip() for t in node.iter(W+'t')) and not any(E.QName(n).localname in ['drawing','pict','sectPr','br'] for n in node.iter())
        if blank:streak.append(node);continue
        if len(streak)>30:
            insertion=body.index(streak[0])
            for old in streak:body.remove(old)
            p=E.Element(W+'p');run=E.SubElement(p,W+'r');E.SubElement(run,W+'br',{W+'type':'page'});body.insert(insertion,p)
        streak=[]
    files['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='utf-8')
    return files

def resolve_template(snapshot):
    manifest=json.loads((ROOT/'assets/templates/manifest.json').read_text(encoding='utf-8'))
    template=next((t for t in manifest['templates'] if t['id']==snapshot['templateId'] and (snapshot.get('templateStorageKey') or str(t['version'])==str(snapshot.get('templateVersion',1)))),None)
    if not template: raise ValueError('TEMPLATE_VERSION_UNKNOWN')
    if snapshot.get('templateStorageKey'):
        store=Path(os.environ['DEMO_ARTIFACT_ROOT']).resolve();path=(store/snapshot['templateStorageKey']).resolve()
        if not path.is_relative_to(store):raise ValueError('TEMPLATE_PATH')
        checksum=snapshot.get('templateChecksum')
    else:
        path=ROOT/'assets/templates'/template['file'];checksum=template['sha256']
    if hashlib.sha256(path.read_bytes()).hexdigest()!=checksum: raise ValueError('TEMPLATE_HASH_MISMATCH')
    return template,path

def preflight(payload,out):
    snapshots=payload.get('snapshots',[])
    if not isinstance(snapshots,list) or len(snapshots)>1000: raise ValueError('PREFLIGHT_LIMIT')
    issues=[]
    for index,snapshot in enumerate(snapshots):
        template,path=resolve_template(snapshot)
        if len(snapshot['items'])!=1: raise ValueError('PREFLIGHT_SINGLE_RECIPIENT')
        try: render_one(snapshot,snapshot['items'][0],path)
        except ValueError as error:
            if str(error)!='PRINT_LAYOUT_OVERFLOW': raise
            issues.append({'index':index,'code':'PRINT_LAYOUT_OVERFLOW'})
    result={'issues':issues,'checked':len(snapshots),'rendererVersion':RENDERER_VERSION}
    Path(out).write_text(json.dumps(result),encoding='utf8');return result

def render_docx(snapshot,out):
    template,path=resolve_template(snapshot)
    items=snapshot['items']
    if not 1<=len(items)<=100: raise ValueError('ROW_LIMIT')
    files=render_one(snapshot,items[0],path)
    root=E.fromstring(files['word/document.xml']); body=root.find(W+'body'); section=body.find(W+'sectPr')
    if section is not None: body.remove(section)
    rels=E.fromstring(files['word/_rels/document.xml.rels'])
    numbering=E.fromstring(files['word/numbering.xml']) if 'word/numbering.xml' in files else None
    for index,item in enumerate(items[1:],1):
        row=render_one(snapshot,item,path); rowroot=E.fromstring(row['word/document.xml']); rowbody=rowroot.find(W+'body')
        rowrels=E.fromstring(row['word/_rels/document.xml.rels'])
        for rel in rowrels:
            if not rel.get('Type','').endswith('/image'): continue
            oldid=rel.get('Id'); new_id=f'rIdPhoto{index}_{oldid}'; oldtarget=rel.get('Target'); newtarget=f'media/batch-{index}-{Path(oldtarget).name}'
            for n in rowroot.iter():
                for attr in [R+'id',R+'embed']:
                    if n.get(attr)==oldid: n.set(attr,new_id)
            copy=deepcopy(rel); copy.set('Id',new_id); copy.set('Target',newtarget); rels.append(copy)
            files['word/'+newtarget]=row['word/'+oldtarget]
        # Section boundary resets the source panel anchors for each recipient.
        p=E.SubElement(body,W+'p');pr=E.SubElement(p,W+'pPr');boundary=deepcopy(section)
        if boundary is not None:
            kind=boundary.find(W+'type')
            if kind is None:kind=E.SubElement(boundary,W+'type')
            kind.set(W+'val','continuous');pr.append(boundary)
        first=rowbody.find(W+'p')
        if first is not None:
            firstpr=first.find(W+'pPr')
            if firstpr is None:firstpr=E.Element(W+'pPr');first.insert(0,firstpr)
            E.SubElement(firstpr,W+'pageBreakBefore')
        for node in rowroot.iter():
            local=E.QName(node).localname
            if local in ['docPr','cNvPr'] and node.get('id'): node.set('id',str(index*1000000+int(node.get('id'))%1000000))
            if local in ['shape','rect','group'] and node.get('id'): node.set('id',node.get('id')+f'-row{index}')
            for attr,value in list(node.attrib.items()):
                key=E.QName(attr).localname
                if key in ['anchorId','editId']:node.set(attr,hashlib.sha256(f'{index}:{value}'.encode()).hexdigest()[:8].upper())
                elif key=='spid' and re.fullmatch(r'_x0000_s\d+',value):node.set(attr,'_x0000_s'+str(index*1000000+int(value.removeprefix('_x0000_s'))))
                elif local in ['bookmarkStart','bookmarkEnd'] and key=='id':node.set(attr,str(index*1000000+int(value)))
                elif local=='bookmarkStart' and key=='name':node.set(attr,value+f'_row{index}')
        if numbering is not None:
            copied={}
            for node in rowroot.iter(W+'numId'):
                old=node.get(W+'val')
                if old=='0':continue
                if old not in copied:
                    definition=next((n for n in numbering.findall(W+'num') if n.get(W+'numId')==old),None)
                    if definition is None:continue
                    new=str(index*10000+int(old));copied[old]=new;duplicate=deepcopy(definition);duplicate.set(W+'numId',new)
                    override=duplicate.find(W+'lvlOverride')
                    if override is None:override=E.SubElement(duplicate,W+'lvlOverride',{W+'ilvl':'0'})
                    start=override.find(W+'startOverride')
                    if start is None:start=E.SubElement(override,W+'startOverride')
                    start.set(W+'val','1');numbering.append(duplicate)
                node.set(W+'val',copied[old])
        for child in rowbody:
            if child.tag!=W+'sectPr': body.append(deepcopy(child))
    if section is not None: body.append(section)
    files['word/document.xml']=E.tostring(root,xml_declaration=True,encoding='utf-8')
    files['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='utf-8')
    if numbering is not None:files['word/numbering.xml']=E.tostring(numbering,xml_declaration=True,encoding='utf-8')
    deterministic_zip(out,normalize_package(files))
    return {'format':'DOCX','templateVersion':snapshot.get('templateVersion',template['version']),'rendererVersion':RENDERER_VERSION}

def convert_pdf(docx,out):
    executable=os.environ.get('DEMO_SOFFICE') or shutil.which('soffice')
    if not executable: raise ValueError('CONVERTER_NOT_INSTALLED')
    # Concurrent batch checks must not race multiple default-profile --version
    # processes. Revalidate whenever the selected executable changes on disk.
    stat=Path(executable).stat();key=(str(Path(executable).resolve()),stat.st_mtime_ns,stat.st_size)
    with _converter_version_lock:
        if key not in _converter_versions:
            version=subprocess.run([executable,'--version'],capture_output=True,timeout=30,text=True).stdout
            if not re.search(r'LibreOffice 26\.2\.6\.3(?:\s|$)',version): raise ValueError('CONVERTER_VERSION_MISMATCH')
            _converter_versions[key]=version
    with tempfile.TemporaryDirectory(prefix='demo-office-') as temp:
        profile=Path(temp)/'profile'; output=Path(temp)/'out'; output.mkdir()
        # The supported 100-recipient cards contain 200 physical pages. A fresh
        # two-converter Linux run exceeded 90s for PS; retain a bounded deadline
        # below the API's PDF-specific 210s process timeout.
        result=subprocess.run([executable,'-env:UserInstallation='+profile.as_uri(),'--headless','--nologo','--nodefault','--nolockcheck','--norestore','--convert-to','pdf:writer_pdf_Export','--outdir',str(output),str(Path(docx).resolve())],capture_output=True,timeout=180)
        pdf=output/(Path(docx).stem+'.pdf')
        if result.returncode or not pdf.exists(): raise ValueError('PDF_CONVERSION_FAILED')
        shutil.copyfile(pdf,out)
    return {'format':'PDF','rendererVersion':RENDERER_VERSION}

def runtime_health(out):
    for package,version in [('Pillow','12.3.0'),('lxml','6.1.1'),('openpyxl','3.1.5'),('defusedxml','0.7.1')]:
        if importlib.metadata.version(package)!=version:raise ValueError('PYTHON_DEPENDENCY_VERSION_MISMATCH')
    manifest=json.loads((ROOT/'assets/templates/manifest.json').read_text(encoding='utf8'))
    for t in manifest['templates']:
        if hashlib.sha256((ROOT/'assets/templates'/t['file']).read_bytes()).hexdigest()!=t['sha256']:raise ValueError('TEMPLATE_HASH_MISMATCH')
    fonts=json.loads((ROOT/'assets/fonts/manifest.json').read_text(encoding='utf8'))
    for filename,checksum in fonts['files'].items():
        if hashlib.sha256((ROOT/'assets/fonts'/filename).read_bytes()).hexdigest()!=checksum:raise ValueError('FONT_HASH_MISMATCH')
        installed=Path(os.environ['LOCALAPPDATA'])/'Microsoft/Windows/Fonts'/filename if os.name=='nt' else Path('/usr/local/share/fonts/demo')/filename
        if not installed.is_file() or hashlib.sha256(installed.read_bytes()).hexdigest()!=checksum:raise ValueError('PINNED_FONTS_NOT_INSTALLED')
    executable=os.environ.get('DEMO_SOFFICE') or shutil.which('soffice')
    if not executable:raise ValueError('CONVERTER_NOT_INSTALLED')
    version=subprocess.run([executable,'--version'],capture_output=True,timeout=15,text=True).stdout
    if not re.search(r'LibreOffice 26\.2\.6\.3(?:\s|$)',version):raise ValueError('CONVERTER_VERSION_MISMATCH')
    result={'ready':True,'rendererVersion':RENDERER_VERSION,'templates':len(manifest['templates']),'fonts':len(fonts['files']),'converter':'26.2.6.3'}
    Path(out).write_text(json.dumps(result),encoding='utf8');return result

COLUMNS=['requestId','status','revision','createdAt','id','fullNameRu','fullNameKz','positionRu','positionKz','workplaceRu','workplaceKz','departmentRu','departmentKz','employerBin','employerAddressRu','employerAddressKz','templateId','direction','documentKind','number','protocolNumber','registrationNumber','documentDate','protocolDate','trainingStart','trainingEnd','trainingSubject','result','reason','education','hours','productionHours','validUntil','externalBasisNumber','biotCategory','biotIndustryRu','biotIndustryKz','biotCheckType','biotKnowledgeResult','biotProctoringResult','biotUniqueNumber','biotNotes']

def export_registry(payload,out):
    wb=Workbook(); ws=wb.active; ws.title='Реестр'; ws.append(COLUMNS)
    for row in payload['items']:
        row={**row,**row.get('assignment',{})}
        row.setdefault('direction',str(row.get('templateId','')).split('-')[0].upper())
        row.setdefault('documentKind','-'.join(str(row.get('templateId','')).split('-')[1:]))
        ws.append([str(row.get(key) or '') for key in COLUMNS])
        for cell in ws[ws.max_row]: cell.data_type='s'; cell.number_format='@'
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    from openpyxl.styles import Font,PatternFill,Alignment
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='163E45')
    for cells in ws.columns:
        ws.column_dimensions[cells[0].column_letter].width=min(46,max(16,max(len(str(c.value or '')) for c in cells)+2))
        for c in cells: c.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(out); return {'rows':len(payload['items']),'columns':COLUMNS,'format':'XLSX'}

def import_table(payload,out):
    path=Path(payload['inputPath']); mapping=payload.get('mapping',{}); errors=[]
    if path.stat().st_size>10*1024*1024: raise ValueError('IMPORT_BYTES_LIMIT')
    suffix=payload.get('format',path.suffix.lstrip('.')).lower(); sheets=[]; selected_sheet=None
    if suffix=='xlsx':
        with ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist())>30*1024*1024 or len(z.infolist())>2000: raise ValueError('IMPORT_ARCHIVE_LIMIT')
            if any('vbaProject' in n or 'externalLinks/' in n for n in z.namelist()): raise ValueError('IMPORT_ACTIVE_CONTENT')
            validate_xlsx_xml(z)
        wb=load_workbook(path,read_only=True,data_only=False,keep_links=False); sheets=wb.sheetnames
        selected_sheet=payload.get('sheet') or sheets[0]
        if selected_sheet not in sheets: raise ValueError('IMPORT_SHEET_UNKNOWN')
        ws=wb[selected_sheet]
        if ws.max_row>10002 or ws.max_column>100: raise ValueError('IMPORT_DIMENSION_LIMIT')
        table=[]
        for rowno,cells in enumerate(ws.iter_rows(),1):
            row=[]
            for cell in cells:
                if cell.data_type=='f': errors.append({'row':rowno,'column':cell.column,'code':'FORMULA_NOT_ALLOWED'}); row.append('')
                elif cell.value is None: row.append('')
                elif isinstance(cell.value,(int,float)) and re.fullmatch('0+',cell.number_format or ''): row.append(str(int(cell.value)).zfill(len(cell.number_format)))
                else: row.append(str(cell.value))
            table.append(row)
        wb.close()
    elif suffix in ['csv','tsv','txt']:
        text=path.read_text(encoding='utf-8-sig')
        delimiter='\t' if suffix in ['tsv','txt'] else payload.get('delimiter')
        if delimiter is None:
            try: delimiter=csv.Sniffer().sniff(text[:65536],delimiters=',;\t').delimiter
            except csv.Error: delimiter=','
        table=list(csv.reader(io.StringIO(text),delimiter=delimiter))
    else: raise ValueError('IMPORT_FORMAT')
    headers=table[0] if table else []; rows=[]; raw_rows=[]; seen={}
    for rowno,row in enumerate(table[1:],2):
        if not any(row) and not any(error.get('row')==rowno for error in errors): continue
        raw_rows.append({'rowNumber':rowno,'values':[str(cell) for cell in row],
                         'errors':[error['code'] for error in errors if error.get('row')==rowno]})
        record={mapping.get(header,header):row[i] if i<len(row) else '' for i,header in enumerate(headers) if mapping.get(header,header) in COLUMNS}
        fingerprint=hashlib.sha256(json.dumps(record,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
        row_errors=[]
        if fingerprint in seen: row_errors.append('DUPLICATE_ROW')
        else: seen[fingerprint]=rowno
        if not record.get('fullNameRu'): row_errors.append('NAME_REQUIRED')
        for key in ['documentDate','protocolDate','trainingStart','trainingEnd','validUntil']:
            if record.get(key):
                try: date.fromisoformat(record[key])
                except ValueError: row_errors.append(key+':INVALID_DATE')
        rows.append({'rowNumber':rowno,'sourceId':fingerprint,'values':record,'errors':row_errors})
    if len(rows)>100: errors.append({'code':'ROW_LIMIT','count':len(rows),'limit':100})
    result={'headers':headers,'sheets':sheets,'sheet':selected_sheet,'rawRows':raw_rows,'rows':rows,'errors':errors,'count':len(rows),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'canApply':len(rows)<=100 and not errors and not any(r['errors'] for r in rows)}
    Path(out).write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8'); return result

def build_bundle(payload,out):
    root=Path(os.environ['DEMO_ARTIFACT_ROOT']).resolve(); entries=[]; missing=list(payload.get('missing',[])); content={}
    for artifact in payload['artifacts']:
        path=(root/artifact['storageKey']).resolve()
        if not path.is_relative_to(root): raise ValueError('ARTIFACT_PATH')
        if not path.is_file(): missing.append({'id':artifact['id'],'reason':'FILE_MISSING'}); continue
        data=path.read_bytes(); digest=hashlib.sha256(data).hexdigest()
        if digest!=artifact['sha256']: missing.append({'id':artifact['id'],'reason':'HASH_MISMATCH'}); continue
        name=re.sub(r'[^\w.\-]','_',artifact.get('fileName') or artifact['id']+'.'+artifact['format'].lower())
        if name in content: name=artifact['id']+'-'+name
        content[name]=data; entries.append({'id':artifact['id'],'file':name,'sha256':digest,'size':len(data),'format':artifact['format']})
    manifest={'version':1,'issuanceId':payload.get('issuanceId'),'complete':not missing and len(entries)==payload.get('expectedCount',len(payload['artifacts'])),'files':entries,'missing':missing,'expectedCount':payload.get('expectedCount',len(payload['artifacts']))}
    content['manifest.json']=json.dumps(manifest,ensure_ascii=False,indent=2).encode()
    content['STATUS.txt']=(('ПОЛНЫЙ КОМПЛЕКТ' if manifest['complete'] else 'НЕПОЛНЫЙ КОМПЛЕКТ')+f"\nФайлов: {len(entries)} из {manifest['expectedCount']}.\nПроверьте manifest.json: состав, контрольные суммы и причины отсутствия файлов.\n").encode('utf8')
    deterministic_zip(out,content)
    return manifest

def main():
    command,out=sys.argv[1:3]; payload=json.loads(sys.stdin.buffer.read().decode('utf8')); Path(out).parent.mkdir(parents=True,exist_ok=True)
    if command=='health': result=runtime_health(out)
    elif command=='preflight': result=preflight(payload,out)
    elif command=='docx': result=render_docx(payload,out)
    elif command=='pdf': result=convert_pdf(payload['docxPath'],out)
    elif command=='photo': result=photo_normalize(payload['inputPath'],out,payload)
    elif command=='import': result=import_table(payload,out)
    elif command=='xlsx': result=export_registry(payload,out)
    elif command=='zip': result=build_bundle(payload,out)
    else: raise ValueError('UNKNOWN_COMMAND')
    if Path(out).exists(): result.update(size=Path(out).stat().st_size,sha256=hashlib.sha256(Path(out).read_bytes()).hexdigest())
    print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':
    try: main()
    except Exception as error:
        # Never include input payload, paths, Python traces or personal data in worker logs.
        print(json.dumps({'error':str(error) if isinstance(error,ValueError) else type(error).__name__}),file=sys.stderr)
        sys.exit(1)

