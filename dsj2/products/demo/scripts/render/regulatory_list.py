"""Private internal training list, using the product's existing LibreOffice converter."""
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_ORIENT

def render_regulatory_list(payload, out, convert_pdf):
    items = payload.get('items', [])
    columns = payload.get('columns', [])
    if not 1 <= len(columns) <= 20 or len(items) > 10000:
        raise ValueError('REGULATORY_LIST_LIMIT')
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.3)
    section.left_margin = section.right_margin = Cm(1.3)
    style = document.styles['Normal']
    style.font.name = 'DejaVu Sans'
    style.font.size = Pt(8)
    document.add_heading(str(payload['title']), 1)
    document.add_paragraph(str(payload['note']))
    table = document.add_table(rows=1, cols=len(columns))
    table.style = 'Table Grid'
    from docx.oxml import OxmlElement
    header = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for cell, column in zip(table.rows[0].cells, columns):
        cell.text = str(column['title'])
    for item in items:
        for cell, column in zip(table.add_row().cells, columns):
            cell.text = str(item.get(column['field'], ''))
    with TemporaryDirectory(prefix='demo-regulatory-') as temp:
        path = Path(temp) / 'internal-list.docx'
        document.save(path)
        convert_pdf(path, out)
    return {'rows': len(items), 'format': 'PDF', 'scope': 'INTERNAL_LIST'}
