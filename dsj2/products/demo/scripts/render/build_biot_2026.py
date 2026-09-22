"""Author current BIOT paper forms from official order223 appendices1/3/4/12.

Development-only builder (python-docx). Runtime reads the pinned DOCX bytes.
No original template is edited in place. --activate archives the old active files.
"""
from pathlib import Path
import argparse,hashlib,io,json,re,shutil
from zipfile import ZipFile
from docx import Document
from docx.shared import Mm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as E
from package_xml import normalize_package
from sanitize_templates import deterministic_zip

ROOT=Path(__file__).resolve().parents[2]
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
OUT=ROOT/'.runtime/biot-2026-source-templates'
VERSIONS={'biot-worker-card':16,'biot-itr-certificate':14,'biot-protocol':10,'biot-itr-protocol':1}

def node(parent,name,**attrs):
 n=OxmlElement('w:'+name)
 for key,value in attrs.items():n.set(qn('w:'+key),str(value))
 parent.append(n);return n

def para(container,value='',size=10,bold=False,center=False,after=3,before=0):
 p=container.add_paragraph();p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
 p.paragraph_format.line_spacing=Pt(size+2);p.paragraph_format.keep_together=True
 if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=p.add_run(value);r.bold=bold;r.font.name='Liberation Serif';r.font.size=Pt(size)
 r._r.get_or_add_rPr().rFonts.set(qn('w:cs'),'Liberation Serif')
 return p

def clear(cell):
 for p in list(cell._tc.findall(qn('w:p'))):cell._tc.remove(p)

def table(doc,widths,rows,caption=None,borders=True):
 t=doc.add_table(rows=rows,cols=len(widths));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 pr=t._tbl.tblPr
 if borders:
  b=node(pr,'tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:node(b,edge,val='single',sz=6,color='333333')
 else:
  b=node(pr,'tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:node(b,edge,val='nil')
 mar=node(pr,'tblCellMar')
 for edge,amount in [('top',70),('left',90),('bottom',70),('right',90)]:node(mar,edge,w=amount,type='dxa')
 # Keep generated table properties in OOXML schema order without changing any
 # historical package bytes in the runtime normalizer.
 if caption:node(pr,'tblCaption',val=caption)
 order='tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange'.split()
 children=list(pr)
 for child in children:pr.remove(child)
 for child in sorted(children,key=lambda n:order.index(E.QName(n).localname) if E.QName(n).localname in order else len(order)):pr.append(child)
 for grid,width in zip(t._tbl.tblGrid,widths):grid.set(qn('w:w'),str(round(width*20)))
 for row in t.rows:
  node(row._tr.get_or_add_trPr(),'cantSplit')
  for cell,width in zip(row.cells,widths):cell.width=Pt(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP;clear(cell)
 return t

def document(landscape=False):
 d=Document();s=d.sections[0];s.page_width=Mm(297 if landscape else 210);s.page_height=Mm(210 if landscape else 297)
 s.top_margin=Mm(13);s.bottom_margin=Mm(13);s.left_margin=Mm(13);s.right_margin=Mm(13)
 s.header_distance=Mm(5);s.footer_distance=Mm(5)
 d.styles['Normal'].font.name='Liberation Serif';d.styles['Normal'].font.size=Pt(10)
 d.styles['Normal'].paragraph_format.space_after=Pt(0)
 for p in list(d._element.body.findall(qn('w:p'))):d._element.body.remove(p)
 return d

def issuer_band(t):
 for cell in t.rows[0].cells:cell.add_paragraph()
 c=t.rows[0].cells[0].merge(t.rows[0].cells[-1]);clear(c);node(c._tc.get_or_add_tcPr(),'shd',fill='E4E7E7',val='clear')
 para(c,'{{ISSUER_RU}}',11,True,True,2);para(c,'{{ISSUER_KZ}}',11,True,True,2)

def committee(c,kz=False,size=10):
 para(c,'Комиссия төрағасы:' if kz else 'Председатель комиссии:',size,True,after=2)
 para(c,'{{CHAIR_POSITION}}',size,after=2);para(c,'{{CHAIR_NAME}}',size,after=2)
 para(c,'Қолы: ____________________' if kz else 'Подпись: ____________________',size,after=6)
 para(c,'Комиссия мүшелері:' if kz else 'Члены комиссии:',size,True,after=2)
 for n in [1,2]:
  para(c,'{{MEMBER_'+str(n)+'_POSITION}}',size,after=2);para(c,'{{MEMBER_'+str(n)+'_NAME}}',size,after=2)
  para(c,'Қолы: ____________________' if kz else 'Подпись: ____________________',size,after=5)

def worker():
 d=document();para(d,'',1,after=0);t=table(d,[260.8,260.8],2,'BIOT2026_APP4');issuer_band(t)
 for cell,kz in zip(t.rows[1].cells,[True,False]):
  para(cell,'КУӘЛІК' if kz else 'УДОСТОВЕРЕНИЕ',14,True,True,6)
  para(cell,'Еңбек қауіпсіздігі және еңбекті қорғау жөніндегі қағидалар, нормалар мен нұсқаулықтар бойынша білімді тексеру емтиханынан өткені туралы' if kz else 'о прохождении экзамена по проверке знаний правил, норм и инструкций по безопасности и охране труда',10,False,True,8)
  para(cell,'{{FULL_NAME_KZ}}' if kz else '{{FULL_NAME_RU}}',11,True,True,3)
  para(cell,'(тегі, аты, әкесінің аты (ол болған жағдайда))' if kz else '(фамилия, имя, отчество (при его наличии))',9,False,True,8)
  para(cell,'Оқу бағдарламасын сәтті аяқтады:' if kz else 'Успешно закончил(а) программу обучения:',10,after=3)
  para(cell,'{{SUBJECT}}',10,True,after=10)
  committee(cell,kz)
  para(cell,('Қала: {{CITY_KZ}}' if kz else 'Город: {{CITY_RU}}'),10,after=5)
  para(cell,('Тіркеу нөмірі: {{REGISTRATION_NUMBER_DISPLAY}}' if kz else 'Рег. № {{REGISTRATION_NUMBER_DISPLAY}}'),10,True,after=5)
  para(cell,'{{DOCUMENT_DATE}}',10,after=7)
  para(cell,'Мөр (МО)' if kz else 'М.П.',10,after=3)
 return d

def certificate():
 d=document(True);para(d,'',1,after=0);t=table(d,[384.1,384.1],2,'BIOT2026_APP1');issuer_band(t)
 for cell,kz in zip(t.rows[1].cells,[True,False]):
  para(cell,'С Е Р Т И Ф И К А Т',16,True,True,8)
  para(cell,('Бірегей нөмірі № {{NUMBER}}' if kz else 'Уникальный номер № {{NUMBER}}'),11,True,True,8)
  para(cell,'берілді' if kz else 'выдан',10,False,True,3)
  para(cell,'{{FULL_NAME_KZ}}' if kz else '{{FULL_NAME_RU}}',13,True,True,3)
  para(cell,'(тегі, аты, әкесінің аты (ол болған жағдайда))' if kz else '(фамилия, имя, отчество (при его наличии))',9,False,True,8)
  value=('Осы арқылы, ол еңбек қауіпсіздігі және еңбекті қорғау саласындағы арнайы кәсіптік құзыреттерді дамыту бойынша оқытудан сәтті өткен және білімді тексеруден өткені {{BIOT_INDUSTRY_KZ}} саласы үшін. оқу бағдарламасына сәйкес ({{HOURS}} сағат) расталады.' if kz else 'В том, что он (она) успешно прошел(шла) обучение и сдал(а) проверку знаний по безопасности и охране труда в соответствии с программой обучения по развитию специальных профессиональных компетенций по вопросам безопасности и охраны труда ({{HOURS}} часов) для {{BIOT_INDUSTRY_RU}} отрасли.')
  para(cell,value,11,after=9)
  para(cell,('Білімді тексеру нәтижесі: {{RESULT}}' if kz else 'Результат проверки знаний: {{RESULT}}'),10,after=5)
  para(cell,('Берілген күні: {{DOCUMENT_DATE}}' if kz else 'Дата выдачи: {{DOCUMENT_DATE}}'),10,after=5)
  para(cell,('Жарамдылық мерзімі: {{VALID_DATE}}' if kz else 'Срок действия: {{VALID_DATE}}'),10,after=5)
  para(cell,('Қала: {{CITY_KZ}}' if kz else 'Город: {{CITY_RU}}'),10,after=8)
  para(cell,'Оқытушы ұйымының басшысының тегі, аты, әкесінің аты (ол бар болған жағдайда):' if kz else 'Фамилия, имя, отчество (при его наличии) руководителя обучающей организации:',10,after=3)
  para(cell,'{{ISSUER_HEAD_NAME}}',11,True,after=6)
  para(cell,'Қолы: ____________________    Мөр (МО)' if kz else 'Подпись: ____________________    М.П.',10,after=3)
 return d

def protocol(special=False):
 d=document(True);caption='BIOT2026_APP12' if special else 'BIOT2026_APP3'
 para(d,'{{ISSUER_RU}} / {{ISSUER_KZ}}',11,True,True,5)
 para(d,('Протокол заседания экзаменационной комиссии по проверке знаний по безопасности и охране труда № {{PROTOCOL_NUMBER}}' if special else 'Протокол заседания экзаменационной комиссии по проверке знаний по безопасности и охране труда (по рабочим профессиям) № {{PROTOCOL_NUMBER}}'),11,True,True,3)
 para(d,('Еңбек қауіпсіздігі және еңбекті қорғау жөніндегі білімді тексеру жөніндегі емтихан комиссиясы отырысының хаттамасы' if special else 'Еңбек қауіпсіздігі және еңбекті қорғау бойынша (жұмыс мамандықтары бойынша) білімді тексеру экзаменациялық комиссиясының отырысының хаттамасы'),10,False,True,5)
 if special:
  para(d,'БИН обучающей организации / Оқытушы ұйымның БСН: {{ISSUER_BIN}}',10,after=3)
  para(d,'Категория обучаемых / Білім алушылар санаты: {{BIOT_CATEGORY_RU}} / {{BIOT_CATEGORY_KZ}}',10,after=3)
  para(d,'Количество академических часов / Оқу бағдарламасының академиялық сағаттарының саны: {{HOURS}}',10,after=3)
 para(d,'Дата проверки знаний / Білімді тексеру күні: {{PROTOCOL_DATE}}',10,after=4)
 para(d,'Комиссия / Комиссия құрамы:',10,True,after=2)
 para(d,'Председатель / Төраға: {{CHAIR_POSITION}} — {{CHAIR_NAME}}',10,after=2)
 para(d,'Члены / Мүшелері: {{MEMBER_1_POSITION}} — {{MEMBER_1_NAME}}; {{MEMBER_2_POSITION}} — {{MEMBER_2_NAME}}',10,after=3)
 para(d,'Основание / Негіздеме: {{APPROVAL_BASIS}}',10,after=3)
 para(d,'Вид проверки знаний / Білімді тексеру түрі: {{BIOT_CHECK_TYPE_RU}} / {{BIOT_CHECK_TYPE_KZ}}',10,after=5)
 if special:
  widths=[24,63,116,105,80,93,65,73,83,66]
  headers=['№ п/п','БИН предприятия\nКәсіпорын БСН','Наименование и адрес предприятия\nКәсіпорын атауы мен мекенжайы','Ф.И.О. (при его наличии) экзаменуемого\nТестіленушінің тегі, аты, әкесінің аты (ол болған жағдайда)','Должность\nЛауазымы','Уникальный номер HS-БИН ОО/номер п/п\nБірегей нөмір ERM/LHS-БСН/курстың аяқталу күні/тіркеу нөмірі','Результат проверки знаний\nБілімді тексеру нәтижесі','Результат прокторинга (прошел/не прошел)\nПрокторинг нәтижесі (өткен/өтпеген)','Отметка о проверке знаний (сдал/не сдал)\nБілімді тексеру белгісі (тапсырды/тапсырмады)','Примечание\nЕскертпе']
  values=['1','{{EMPLOYER_BIN}}','{{EMPLOYER_NAME_ADDRESS_BOTH}}','{{FULL_NAME_BOTH}}','{{POSITION_BOTH}}','{{BIOT_UNIQUE_NUMBER}}','{{BIOT_KNOWLEDGE_RESULT}}','{{BIOT_PROCTORING_RESULT}}','{{RESULT}}','{{BIOT_NOTES}}']
 else:
  widths=[27,129,148,106,110,170,78]
  headers=['№ п/п','Наименование организации\nҰйымның атауы','Ф.И.О. (при его наличии) экзаменуемого\nТестіленушінің тегі, аты, әкесінің аты (ол болған жағдайда)','Профессия (должность)\nМамандық (лауазым)','Участок, цех, подразделение\nУчасток, цех, бөлімше','Отметка о проверке знаний (сдал, подлежит повторной проверке знаний по безопасности и охране труда)\nБілімді тексеру белгісі (тапсырған / білімді қайтадан тексеруге тиіс)','Примечание\nЕскертпе']
  values=['1','{{WORKPLACE_BOTH}}','{{FULL_NAME_BOTH}}','{{POSITION_BOTH}}','{{DEPARTMENT_BOTH}}','{{RESULT}}','{{BIOT_NOTES}}']
 total=sum(widths);widths=[w*768.2/total for w in widths]
 t=table(d,widths,3,caption)
 for i,row in enumerate(t.rows):
  if i<2:node(row._tr.get_or_add_trPr(),'tblHeader')
  for j,cell in enumerate(row.cells):
   value=headers[j] if i==0 else str(j+1) if i==1 else values[j]
   p=para(cell,value,9,False,i==1,after=0)
   if i==1:
    pr=p._p.get_or_add_pPr();n=node(pr,'numPr');node(n,'numId',val=0)
 para(d,'',9,after=3)
 para(d,'Председатель комиссии / Комиссия төрағасы: {{CHAIR_NAME}}    ____________________',10,after=4)
 para(d,'Члены комиссии / Комиссия мүшелері: {{MEMBER_1_NAME}}    ____________________',10,after=4)
 para(d,'{{MEMBER_2_NAME}}    ____________________',10,after=3)
 if not special:para(d,'М.П. / Мөр (МО)',10,after=0)
 return d

def build(path,doc):
 buffer=io.BytesIO();doc.save(buffer)
 with ZipFile(buffer) as z:files={n:z.read(n) for n in z.namelist() if not n.startswith('docProps/thumbnail')}
 deterministic_zip(path,normalize_package(files))

def main():
 parser=argparse.ArgumentParser();parser.add_argument('--activate',action='store_true');args=parser.parse_args()
 OUT.mkdir(parents=True,exist_ok=True);catalog=ROOT/'assets/templates';manifest=json.loads((catalog/'manifest.json').read_text(encoding='utf8'))
 builders={'biot-worker-card':worker,'biot-itr-certificate':certificate,'biot-protocol':lambda:protocol(False),'biot-itr-protocol':lambda:protocol(True)}
 changes=[]
 for tid,builder in builders.items():
  path=OUT/(tid+f'.v{VERSIONS[tid]}.docx');build(path,builder());digest=hashlib.sha256(path.read_bytes()).hexdigest()
  with ZipFile(path) as z:
   tree=E.fromstring(z.read('word/document.xml'));fields=sorted(set(re.findall(r'\{\{([A-Z0-9_]+)\}\}',''.join(tree.itertext()))));sections=[dict(n.attrib) for n in tree.iter(W+'pgSz')]
  current=next((t for t in manifest['templates'] if t['id']==tid),None)
  if args.activate and current and current['file']==path.name and current['sha256']!=digest:
   raise ValueError('IMMUTABLE_TEMPLATE_VERSION_COLLISION: increment the affected VERSIONS entry')
  changes.append({'id':tid,'version':VERSIONS[tid],'file':path.name,'sha256':digest,'fields':fields,'sections':sections,'formRevision':'BIOT_2026_223','sourceSha256':'864c4ceac4ceb06e3bb385f229491da2ab2366e91c032424807991dbfed07c3d','verificationStatus':'NOT_RUN'})
 if args.activate:
  historical=ROOT/'docs/evidence/commercial-acceptance/printing/historical-inputs';historical.mkdir(exist_ok=True)
  old_manifest=historical/'manifest-before-biot-2026.json'
  if not old_manifest.exists():shutil.copyfile(catalog/'manifest.json',old_manifest)
  for change in changes:
   current=next((t for t in manifest['templates'] if t['id']==change['id']),None)
   if current:
    previous=catalog/current['file'];target=historical/previous.name
    if previous.name!=change['file']:
     if target.exists():assert target.read_bytes()==previous.read_bytes();previous.unlink()
     else:previous.rename(target)
    change['previousTemplateSha256']=current['sha256']
   else:current={'languages':['ru','kk'],'photo':False,'exports':['DOCX','PDF'],'sourceImagesRemoved':True};manifest['templates'].append(current)
   current.update(change);current.pop('sample',None)
   current.update(legalApproval='REQUIRED_BY_ISSUER',protocolSemantics='INDIVIDUAL' if change['id'].endswith('protocol') else None,regulatoryReview={'status':'CURRENT_PAPER_FORM_UNVERIFIED','effectiveFrom':'2026-07-12','source':'https://zan.gov.kz/api/documents/225864/rus/download/pdf','notice':'Бумажная форма по приказу №223. Геометрия и совместимость новой версии ещё не подтверждены. Выпуск требует правильной категории, фактических результатов и утверждения эмитента; документ не заменяет оригинал ЕЦС.'})
   shutil.copyfile(OUT/change['file'],catalog/change['file'])
  manifest['rendererVersion']='demo-ooxml-6';(catalog/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
  (catalog/'biot-2026-changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
 (OUT/'manifest.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2)+'\n',encoding='utf8');print(json.dumps(changes,ensure_ascii=False))

if __name__=='__main__':main()
