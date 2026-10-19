"""Bounded-memory XLSX output for the trusted server-created JSONL spool."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment


def export_lms_rows(payload, out):
    columns = payload['columns']
    if not 1 <= len(columns) <= 100:
        raise ValueError('EXPORT_COLUMNS_LIMIT')
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('Реестр')
    ws.freeze_panes = 'A2'
    for index in range(len(columns)):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(index + 1)].width = 26

    def cells(values, header=False):
        for value in values:
            cell = WriteOnlyCell(ws, value=str(value if value is not None else ''))
            # Explicit string type prevents formulas, external links, and numeric ID conversion.
            cell.data_type = 's'
            cell.number_format = '@'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if header:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='163E45')
            yield cell

    ws.append(list(cells([column['title'] for column in columns], True)))
    count = 0
    with Path(payload['itemsFile']).open(encoding='utf8') as source:
        for line in source:
            row = json.loads(line)
            ws.append(list(cells([row.get(column['field'], '') for column in columns])))
            count += 1
    if payload.get('snapshotAt'):
        meta = wb.create_sheet('Срез' if payload.get('language') != 'kz' else 'Кесінді')
        for key, value in [('Отчёт / Есеп', payload.get('title')), ('Определение / Анықтама', payload.get('definition')), ('Срез / Кесінді', payload['snapshotAt']), ('Часовой пояс / Уақыт белдеуі', payload.get('timezone')), ('Строк / Жолдар', count), ('Фильтры / Сүзгілер', json.dumps(payload.get('filters', {}), ensure_ascii=False)), ('Статус / Күйі', 'Отчётный срез, не юридический документ / Есептік кесінді, заңды құжат емес')]:
            meta.append(list(cells([key, value])))
        for key, value in payload.get('summary', {}).items():
            from lms_report import SUMMARY_LABELS
            label = SUMMARY_LABELS.get(key, (key, key))[1 if payload.get('language') == 'kz' else 0]
            meta.append(list(cells([label, value])))
        meta.column_dimensions['A'].width = 32
        meta.column_dimensions['B'].width = 85
    wb.save(out)
    return {'rows': count, 'columns': [column['title'] for column in columns], 'format': 'XLSX'}
