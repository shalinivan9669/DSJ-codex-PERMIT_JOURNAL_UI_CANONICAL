import hashlib
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.shared import Inches
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'scripts/render'))
from print_pack import merge_pdf, merge_docx, PART_BYTES
from renderer import convert_pdf, render_docx

class PrintPackTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
    def tearDown(self):
        self.temp.cleanup()
    def artifact(self, path, number, template='form'):
        return {'id': str(number), 'storageKey': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'templateId': template, 'templateVersion': 'v1'}
    def test_pdf_original_page_boxes_rotations_and_order(self):
        files = []
        for index, (width, height, rotation) in enumerate([(300, 500, 0), (842, 595, 90), (200, 250, 0)]):
            path = self.root / ('%d.pdf' % index)
            writer = PdfWriter()
            writer.add_blank_page(width, height).rotate(rotation)
            writer.write(path)
            files.append(self.artifact(path, index))
        out = self.root / 'merged.pdf'
        before = [a['sha256'] for a in files]
        result = merge_pdf({'artifacts': files, 'artifactRoot': str(self.root)}, out)
        pages = PdfReader(out).pages
        self.assertEqual(result['pages'], 3)
        self.assertEqual([(float(p.mediabox.width), float(p.mediabox.height), p.rotation) for p in pages], [(300, 500, 0), (842, 595, 90), (200, 250, 0)])
        self.assertEqual([hashlib.sha256((self.root / a['storageKey']).read_bytes()).hexdigest() for a in files], before)
    def test_docx_saved_text_photo_sections_and_pdf_conversion(self):
        files = []
        for index, color in enumerate(['red', 'blue']):
            image = self.root / ('photo%d.png' % index)
            Image.new('RGB', (80, 120), color).save(image)
            doc = Document()
            doc.add_paragraph('Синтетический Получатель %d — Ә Ғ Қ Ң Ө Ұ Ү Һ І' % index)
            doc.add_paragraph('№ DOC-%d · 04.10.2026' % index)
            doc.add_picture(str(image), width=Inches(.5))
            path = self.root / ('%d.docx' % index)
            doc.save(path)
            files.append(self.artifact(path, index))
        out = self.root / 'merged.docx'
        merge_docx({'artifacts': files, 'artifactRoot': str(self.root)}, out)
        merged = Document(out)
        self.assertEqual(len(merged.sections), 2)
        text = '\n'.join(p.text for p in merged.paragraphs)
        self.assertIn('Получатель 0', text)
        self.assertIn('Получатель 1', text)
        self.assertIn('Ә Ғ Қ Ң Ө Ұ Ү Һ І', text)
        with ZipFile(out) as z:
            pictures = [n for n in z.namelist() if n.endswith('.png')]
            self.assertEqual(len(pictures), 2)
            self.assertNotEqual(z.read(pictures[0]), z.read(pictures[1]))
        pdf = self.root / 'merged.pdf'
        convert_pdf(str(out), str(pdf))
        self.assertEqual(len(PdfReader(pdf).pages), 2)
        for index, page in enumerate(PdfReader(pdf).pages):
            self.assertIn('Получатель %d' % index, page.extract_text())
    def test_rejects_corrupt_sources_empty_sets_and_mixed_docx(self):
        with self.assertRaisesRegex(ValueError, 'EMPTY_PRINT_SET'):
            merge_pdf({'artifacts': [], 'artifactRoot': str(self.root)}, self.root / 'none.pdf')
        path = self.root / 'a.docx'
        Document().save(path)
        first = self.artifact(path, 1)
        second = self.artifact(path, 2, 'different')
        with self.assertRaisesRegex(ValueError, 'DOCX_MIXED_FORMS'):
            merge_docx({'artifacts': [first, second], 'artifactRoot': str(self.root)}, self.root / 'none.docx')
        path.write_bytes(b'corrupt')
        with self.assertRaisesRegex(ValueError, 'ARTIFACT_HASH_MISMATCH'):
            merge_docx({'artifacts': [first], 'artifactRoot': str(self.root)}, self.root / 'none.docx')
    def test_actual_ptm_continuous_sections_and_page_anchors_start_each_saved_card_on_its_own_page(self):
        from test_render import fixture
        sources = []
        for index in range(2):
            snapshot = fixture('ptm-card')
            person = snapshot['items'][0]
            person['fullNameRu'] = 'Синтетический Получатель %d' % index
            person['fullNameKz'] = 'Сынақ Қатысушы %d' % index
            person['number'] = 'ПТМ-TEST-%d' % index
            person['assignment']['outcome'] = {'status': 'PASSED', 'source': 'СИНТЕТИЧЕСКАЯ ведомость проверки страниц'}
            path = self.root / ('saved-ptm-%d.docx' % index)
            render_docx(snapshot, path)
            source = self.artifact(path, index, 'ptm-card')
            source['templateVersion'] = str(snapshot['templateVersion'])
            sources.append(source)
        before = [source['sha256'] for source in sources]
        combined = self.root / 'saved-two-ptm-cards.docx'
        merge_docx({'artifacts': sources, 'artifactRoot': str(self.root)}, combined)
        pdf = self.root / 'saved-two-ptm-cards.pdf'
        convert_pdf(combined, pdf)
        pages = PdfReader(pdf).pages
        self.assertEqual(len(pages), 2, 'Actual floating frames must not overlay both saved people on one page')
        for index, page in enumerate(pages):
            text = page.extract_text()
            self.assertIn('Синтетический Получатель %d' % index, text)
            self.assertIn('ПТМ-TEST-%d' % index, text)
            self.assertNotIn('Получатель %d' % (1 - index), text)
        self.assertEqual([hashlib.sha256((self.root / source['storageKey']).read_bytes()).hexdigest() for source in sources], before)
    def test_source_byte_budget_before_merge(self):
        path = self.root / 'large.pdf'
        path.write_bytes(b'x' * (PART_BYTES + 1))
        with self.assertRaisesRegex(ValueError, 'PRINT_SET_PART_LIMIT'):
            merge_pdf({'artifacts': [self.artifact(path, 1)], 'artifactRoot': str(self.root)}, self.root / 'none.pdf')

if __name__ == '__main__':
    unittest.main()
