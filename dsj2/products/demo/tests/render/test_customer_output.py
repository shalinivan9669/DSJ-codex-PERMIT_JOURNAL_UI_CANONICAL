import hashlib
import json
import os
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'scripts/render'))
from renderer import export_registry, import_table, build_bundle
from openpyxl import Workbook, load_workbook


class CustomerOutputTests(unittest.TestCase):
    def test_xlsx_literal_formulas_header_and_text_identifiers(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'registry.xlsx'
            export_registry({'columns': [{'field': 'tab', 'title': '=1+1'}, {'field': 'name', 'title': 'ФИО'}],
                             'items': [{'tab': '000001', 'name': '=HYPERLINK("bad")'}]}, path)
            wb = load_workbook(path, data_only=False)
            self.assertEqual(wb.active['A1'].data_type, 's')
            self.assertEqual(wb.active['A2'].value, '000001')
            self.assertEqual(wb.active['B2'].data_type, 's')
            self.assertEqual(wb.active['B2'].value, '=HYPERLINK("bad")')
            wb.close()

    def test_excel_date_only_and_zero_mask_survive_import(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'people.xlsx'
            wb = Workbook(); ws = wb.active
            ws.append(['fullNameRu', 'personnelNumber', 'documentDate'])
            ws.append(['Әділ Ғалым', 7, datetime(2026, 9, 22)])
            ws['B2'].number_format = '000000'
            wb.save(path); wb.close()
            parsed = import_table({'inputPath': str(path), 'format': 'xlsx'}, Path(temp) / 'result.json')
            self.assertEqual(parsed['rawRows'][0]['values'], ['Әділ Ғалым', '000007', '2026-09-22'])
            self.assertEqual(parsed['rows'][0]['errors'], [])

    def test_handover_uses_immutable_bytes_inventory_and_safe_paths(self):
        with tempfile.TemporaryDirectory() as temp, patch.dict(os.environ, {'DEMO_ARTIFACT_ROOT': temp}):
            data = b'Synthetic immutable original PDF'
            (Path(temp) / 'original.pdf').write_bytes(data)
            digest = hashlib.sha256(data).hexdigest()
            artifact = {'id': 'one', 'storageKey': 'original.pdf', 'sha256': digest, 'format': 'PDF',
                        'fileName': '../CON/000007 — Әділ Ғалым.pdf', 'provenance': 'RECONSTRUCTED'}
            payload = {'artifacts': [artifact], 'issuanceId': 'one', 'expectedCount': 1,
                       'coverText': 'Передаём сохранённый документ.', 'includeInventory': True}
            result = build_bundle(payload, Path(temp) / 'handover.zip')
            self.assertTrue(result['complete'])
            with ZipFile(Path(temp) / 'handover.zip') as archive:
                self.assertIn('_CON/000007 — Әділ Ғалым.pdf', archive.namelist())
                self.assertEqual(archive.read(result['files'][0]['file']), data)
                self.assertIn('Опись.tsv', archive.namelist())
                self.assertIn('Сопроводительное письмо.txt', archive.namelist())
                manifest = json.loads(archive.read('manifest.json'))
                self.assertEqual(manifest['files'][0]['sha256'], digest)
                self.assertEqual(manifest['files'][0]['provenance'], 'RECONSTRUCTED')
                self.assertEqual(len(manifest['attachments']), 2)
                for derivative in manifest['attachments']:
                    self.assertEqual(hashlib.sha256(archive.read(derivative['file'])).hexdigest(), derivative['sha256'])
            self.assertEqual((Path(temp) / 'original.pdf').read_bytes(), data)
            self.assertEqual(artifact['fileName'], '../CON/000007 — Әділ Ғалым.pdf')
            kazakh = build_bundle({**payload, 'profile': {'language': 'kz'}, 'coverText': 'Сақталған құжаттар жинағы.'}, Path(temp) / 'kazakh.zip')
            with ZipFile(Path(temp) / 'kazakh.zip') as archive:
                self.assertIn('Күйі', archive.read('Тізімдеме.tsv').decode('utf8'))
                self.assertIn('Сақталған', archive.read('Ілеспе хат.txt').decode('utf8'))
                self.assertIn('ТОЛЫҚ ЖИНАҚ', archive.read('STATUS.txt').decode('utf8'))
                self.assertEqual(archive.read(kazakh['files'][0]['file']), data)
            with self.assertRaisesRegex(ValueError, 'EMPTY_BUNDLE'):
                build_bundle({'artifacts': [], 'expectedCount': 0}, Path(temp) / 'empty.zip')

    def test_partial_handover_identifies_missing_and_corrupt_files(self):
        with tempfile.TemporaryDirectory() as temp, patch.dict(os.environ, {'DEMO_ARTIFACT_ROOT': temp}):
            data = b'An existing immutable saved document'
            (Path(temp) / 'ready.pdf').write_bytes(data)
            (Path(temp) / 'corrupt.pdf').write_bytes(b'changed bytes')
            digest = hashlib.sha256(data).hexdigest()
            artifacts = [
                {'id': key, 'storageKey': key + '.pdf', 'sha256': digest,
                 'format': 'PDF', 'fileName': key + '.pdf', 'provenance': 'ORIGINAL'}
                for key in ['ready', 'missing', 'corrupt']
            ]
            result = build_bundle({'artifacts': artifacts, 'expectedCount': 3,
                                   'includeInventory': True,
                                   'coverText': 'Передаём только фактически готовые файлы.'},
                                  Path(temp) / 'partial.zip')
            self.assertFalse(result['complete'])
            self.assertEqual(result['readyCount'], 1)
            self.assertEqual(result['expectedCount'], 3)
            self.assertEqual({item['id']: item['reason'] for item in result['missing']},
                             {'missing': 'FILE_MISSING', 'corrupt': 'HASH_MISMATCH'})
            with ZipFile(Path(temp) / 'partial.zip') as archive:
                self.assertEqual(archive.read('ready.pdf'), data)
                self.assertNotIn('missing.pdf', archive.namelist())
                self.assertNotIn('corrupt.pdf', archive.namelist())
                self.assertIn('НЕПОЛНЫЙ КОМПЛЕКТ', archive.read('STATUS.txt').decode('utf8'))
                cover = archive.read('Сопроводительное письмо.txt').decode('utf8')
                self.assertIn('НЕПОЛНЫЙ КОМПЛЕКТ', cover)
                self.assertIn('Готовых файлов: 1 из 3.', cover)
                self.assertIn('Готов: ready.pdf', cover)
                self.assertIn('Файл отсутствует', cover)
                self.assertIn('Целостность файла не подтверждена', cover)
                inventory = archive.read('Опись.tsv').decode('utf8')
                self.assertIn('FILE_MISSING', inventory)
                self.assertIn('HASH_MISMATCH', inventory)
                manifest = json.loads(archive.read('manifest.json'))
                self.assertFalse(manifest['complete'])
                self.assertEqual(len(manifest['files']), 1)
            self.assertEqual((Path(temp) / 'ready.pdf').read_bytes(), data)


if __name__ == '__main__':
    unittest.main()
