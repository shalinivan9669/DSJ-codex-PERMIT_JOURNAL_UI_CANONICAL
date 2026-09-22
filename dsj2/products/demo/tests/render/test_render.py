"""Mandatory real file verification; missing converter is a failure, never a skip."""
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sys
import unittest
from zipfile import ZipFile
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'scripts/render'))
from renderer import render_docx,convert_pdf,photo_normalize,export_registry,import_table,build_bundle
from openpyxl import Workbook,load_workbook
from PIL import Image
from lxml import etree as E

OUT=ROOT/'docs/evidence/render'
OUT.mkdir(parents=True,exist_ok=True)
os.environ.setdefault('DEMO_ARTIFACT_ROOT',str(OUT/'store'))
STORE=Path(os.environ['DEMO_ARTIFACT_ROOT']);STORE.mkdir(parents=True,exist_ok=True)
MANIFEST=json.loads((ROOT/'assets/templates/manifest.json').read_text(encoding='utf8'))

def fixture(template_id):
    snap={'mode':'issued-document','demoMode':True,'templateId':template_id,'templateVersion':next(t['version'] for t in MANIFEST['templates'] if t['id']==template_id),
       'issuer':{'nameRu':'Учебный центр «Образец»','nameKz':'«Үлгі» оқу орталығы','cityRu':'Астана','cityKz':'Астана','approvalBasis':'ТЕСТ: приказ комиссии № 7 от 10.09.2026','commission':[{'name':'Тестова А. Б.','position':'Председатель'},{'name':'Үлгі Ә. Ө.','position':'Член комиссии'},{'name':'Примеров В. Г.','position':'Член комиссии'}]},
       'items':[{'id':'recipient-01','fullNameRu':'Тестов Иван Петрович','fullNameKz':'Әділбек Өмірсерік Қанатұлы','positionRu':'Инженер','positionKz':'Инженер','workplaceRu':'ТОО «Тестовый заказчик»','workplaceKz':'«Сынақ тапсырыс беруші» ЖШС','number':'ТЕСТ-00001','protocolNumber':'ПР-00001','registrationNumber':'00123','assignment':{'id':'a1','templateId':template_id,'documentDate':'2026-09-22','protocolDate':'2026-09-20','trainingStart':'2026-09-14','trainingEnd':'2026-09-19','trainingSubject':'Безопасность труда / Еңбек қауіпсіздігі','result':'ТЕСТ: хорошо / жақсы','reason':'ТЕСТ: первичное обучение','hours':40,'validUntil':'2027-09-22','externalBasisNumber':'ВНЕШНИЙ-01','protocolMode':'individual'}}], 'photos':{}}
    if template_id.startswith('biot-'):
        snap['issuer'].update(bin='990140000000',headName='Тестова А. Б.')
        item=snap['items'][0];item.update(departmentRu='Испытательный участок',departmentKz='Сынақ учаскесі',employerBin='990240000000',employerAddressRu='Астана, Тестовая 1',employerAddressKz='Астана, Сынақ 1',credentialNumber='СЕРТ-00001')
        item['assignment'].update(biotCategory='OHS_SPECIALIST_SPECIAL' if template_id.startswith('biot-itr-') else 'WORKER',productionHours='16',biotIndustryRu='промышленной',biotIndustryKz='өнеркәсіп',biotCheckType='PERIODIC',biotKnowledgeResult='90 / 100',biotProctoringResult='прошел / өткен',biotUniqueNumber='',biotNotes='')
    return snap

class RenderTests(unittest.TestCase):
    def test_01_all_active_forms_real_docx_pdf(self):
        results=[]
        for template in MANIFEST['templates']:
            with self.subTest(template=template['id']):
                snap=fixture(template['id'])
                if template['photo']:
                    portrait=Image.new('RGB',(480,640),(70,115,150));portrait.save(STORE/'test-photo.png')
                    snap['items'][0]['photoAssetId']='synthetic-photo';snap['photos']['synthetic-photo']='test-photo.png'
                out=OUT/(template['id']+'.docx');render_docx(snap,out)
                with ZipFile(out) as z:
                    xml=b'\n'.join(z.read(n) for n in z.namelist() if n.endswith('.xml')).decode('utf8')
                    self.assertNotRegex(xml,r'\{\{[A-Z_]+\}\}|MERGEFIELD|Стандарт|Солтанова|Флеглер|Баянов|Жакибеков|Есен Д\.')
                    self.assertIn('Тестов',xml)
                    self.assertTrue(all(b'TargetMode="External"' not in z.read(n) for n in z.namelist() if n.endswith('.rels')))
                pdf=out.with_suffix('.pdf');convert_pdf(out,pdf)
                self.assertTrue(pdf.read_bytes().startswith(b'%PDF-'));self.assertGreater(pdf.stat().st_size,2000)
                results.append({'templateId':template['id'],'docxSha256':hashlib.sha256(out.read_bytes()).hexdigest(),'pdfSha256':hashlib.sha256(pdf.read_bytes()).hexdigest()})
        (OUT/'files.json').write_text(json.dumps(results,indent=2),encoding='utf8')
    def test_02_batch_and_long_bilingual(self):
        snap=fixture('biot-worker-card');second=json.loads(json.dumps(snap['items'][0]));second['fullNameRu']='Тестов-Примеров Александр Константинович';second['fullNameKz']='Әділбек Өмірсерік Қанатұлы';second['number']='ТЕСТ-00002';snap['items'].append(second)
        render_docx(snap,OUT/'biot-batch-two.docx');convert_pdf(OUT/'biot-batch-two.docx',OUT/'biot-batch-two.pdf')
        hundred=fixture('biot-protocol');source=hundred['items'][0];hundred['items']=[]
        for index in range(100):
            item=json.loads(json.dumps(source));item.update(id=f'recipient-{index+1}',fullNameRu=f'Тестов-{index+1:03} Иван',fullNameKz=f'Әділбек-{index+1:03} Қанатұлы',number=f'ТЕСТ-{index+1:05}',protocolNumber=f'ПР-{index+1:05}')
            hundred['items'].append(item)
        render_docx(hundred,OUT/'biot-100.docx');convert_pdf(OUT/'biot-100.docx',OUT/'biot-100.pdf')
        from print_contracts import assert_docx_columns,assert_pdf_columns
        assert_docx_columns(OUT/'biot-100.docx','biot-protocol',100)
        assert_pdf_columns(OUT/'biot-100.pdf',100,100)
        hundred['items'].append(hundred['items'][0])
        with self.assertRaisesRegex(ValueError,'ROW_LIMIT'):render_docx(hundred,OUT/'invalid.docx')
    def test_03_safe_registry_all_fields_and_import(self):
        item=fixture('ps-witness')['items'][0];item['fullNameRu']='=1+1';item['registrationNumber']='00012'
        export_registry({'items':[item]},OUT/'registry.xlsx');wb=load_workbook(OUT/'registry.xlsx');ws=wb.active
        headings=[c.value for c in ws[1]];self.assertIn('fullNameKz',headings);self.assertIn('registrationNumber',headings)
        c=ws.cell(2,headings.index('fullNameRu')+1);self.assertEqual(c.data_type,'s');self.assertEqual(c.value,'=1+1');wb.close()
        result=import_table({'inputPath':str(OUT/'registry.xlsx'),'format':'xlsx'},OUT/'import.json')
        self.assertEqual(result['rows'][0]['values']['registrationNumber'],'00012');self.assertTrue(result['canApply'])
    def test_04_formulas_rejected_and_limits(self):
        wb=Workbook();ws=wb.active;ws.append(['fullNameRu','number']);ws.append(['=1+1','001']);wb.save(OUT/'formula-input.xlsx')
        result=import_table({'inputPath':str(OUT/'formula-input.xlsx')},OUT/'formula-result.json');self.assertFalse(result['canApply']);self.assertEqual(result['errors'][0]['code'],'FORMULA_NOT_ALLOWED')
        for count in [100,101]:
            path=OUT/f'{count}.csv';path.write_text('fullNameRu,fullNameKz\n'+'\n'.join(f'Тестов {i},Әділбек {i}' for i in range(count)),encoding='utf8')
            result=import_table({'inputPath':str(path)},OUT/f'{count}.json');self.assertEqual(result['canApply'],count==100)
        path=OUT/'duplicate.csv';path.write_text('fullNameRu\nА\nА\n',encoding='utf8');result=import_table({'inputPath':str(path)},OUT/'duplicate.json');self.assertEqual(result['rows'][1]['errors'],['DUPLICATE_ROW'])
    def test_05_photo_decode_crop_rotation(self):
        path=OUT/'source-photo.png';Image.new('RGB',(800,1000),(100,160,180)).save(path)
        result=photo_normalize(path,OUT/'normalized-photo.png',{'rotation':90,'crop':{'x':0.1,'y':0.1,'width':0.6,'height':0.7}})
        self.assertEqual((result['width'],result['height']),(600,560))
        (OUT/'bad.png').write_bytes(b'not an image')
        with self.assertRaises(Exception):photo_normalize(OUT/'bad.png',OUT/'bad-out.png',{})
    def test_06_bundle_uses_stored_bytes_and_reports_partial(self):
        file=STORE/'saved.docx';file.write_bytes((OUT/'ps-witness.docx').read_bytes());digest=hashlib.sha256(file.read_bytes()).hexdigest()
        artifact={'id':'sample','storageKey':'saved.docx','sha256':digest,'format':'DOCX'}
        manifest=build_bundle({'artifacts':[artifact],'issuanceId':'synthetic','expectedCount':1},OUT/'complete.zip');self.assertTrue(manifest['complete'])
        with ZipFile(OUT/'complete.zip') as z:self.assertEqual(z.read('sample.docx'),file.read_bytes())
        manifest=build_bundle({'artifacts':[artifact],'expectedCount':2},OUT/'partial.zip');self.assertFalse(manifest['complete'])
    def test_07_clean_templates_keep_exact_geometry(self):
        for template in MANIFEST['templates']:
            with ZipFile(ROOT/'assets/templates'/template['file']) as z:
                self.assertFalse(any(n.startswith('word/media/') for n in z.namelist()))
                tree=E.fromstring(z.read('word/document.xml'))
                self.assertEqual([dict(n.attrib) for n in tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgSz')],template['sections'])
    def test_08_no_demo_mark_on_working_documents(self):
        snap=fixture('ps-witness');snap['demoMode']=False;render_docx(snap,OUT/'working-unmarked.docx')
        with ZipFile(OUT/'working-unmarked.docx') as z:self.assertNotIn('word/demo-header.xml',z.namelist())
    def test_09_raw_headers_sheet_choice_and_leading_zeros(self):
        wb=Workbook();wb.active.title='Первый';wb.active.append(['ФИО RU','Код']);wb.active.append(['Первый','0001'])
        ws=wb.create_sheet('Қазақша');ws.append(['ФИО RU','Код']);ws.append(['Әғқңөұүһі',12]);ws['B2'].number_format='00000';wb.save(OUT/'sheets.xlsx')
        result=import_table({'inputPath':str(OUT/'sheets.xlsx'),'sheet':'Қазақша'},OUT/'sheets-import.json')
        self.assertEqual(result['sheets'],['Первый','Қазақша']);self.assertEqual(result['sheet'],'Қазақша')
        self.assertEqual(result['rawRows'][0],{'rowNumber':2,'values':['Әғқңөұүһі','00012'],'errors':[]})
        with self.assertRaisesRegex(ValueError,'IMPORT_SHEET_UNKNOWN'):import_table({'inputPath':str(OUT/'sheets.xlsx'),'sheet':'missing'},OUT/'wrong-sheet.json')
    def test_10_snapshot_template_bytes_survive_manifest_version_change(self):
        snap=fixture('biot-worker-card');template=next(t for t in MANIFEST['templates'] if t['id']==snap['templateId'])
        original=(ROOT/'assets/templates'/template['file']).read_bytes();(STORE/'snapshot-template.docx').write_bytes(original)
        snap.update(templateStorageKey='snapshot-template.docx',templateChecksum=hashlib.sha256(original).hexdigest(),templateVersion='archived-v77')
        first=render_docx(snap,OUT/'snapshot-first.docx');render_docx(snap,OUT/'snapshot-second.docx')
        self.assertEqual(first['templateVersion'],'archived-v77');self.assertEqual((OUT/'snapshot-first.docx').read_bytes(),(OUT/'snapshot-second.docx').read_bytes())
        (STORE/'snapshot-template.docx').write_bytes(original+b'corruption')
        with self.assertRaisesRegex(ValueError,'TEMPLATE_HASH_MISMATCH'):render_docx(snap,OUT/'snapshot-invalid.docx')
    def test_11_hidden_drawing_caches_and_calendar_dates(self):
        from renderer import date_parts
        self.assertEqual(date_parts('2028-02-29')['DATE'],'29.02.2028')
        self.assertEqual(date_parts('2026-12-31')['DATE'],'31.12.2026')
        self.assertEqual(date_parts('2027-01-01')['DATE'],'01.01.2027')
        with self.assertRaises(ValueError):date_parts('2026-02-29')
        for template in MANIFEST['templates']:
            with ZipFile(ROOT/'assets/templates'/template['file']) as z:
                for filename in z.namelist():
                    if filename.endswith('.xml'):self.assertNotIn(b'gfxdata=',z.read(filename))
    def test_13_unbreakable_text_rejected_before_clipping(self):
        snapshot=fixture('pb-card');snapshot['items'][0]['fullNameRu']='А'*500
        with self.assertRaisesRegex(ValueError,'PRINT_LAYOUT_OVERFLOW'):render_docx(snapshot,OUT/'overflow.docx')

    def test_14_independent_bilingual_values_in_single_cells(self):
        for tid in ['biot-protocol','ptm-protocol','pb-protocol','ps-protocol','biot-itr-certificate']:
            snap=fixture(tid);item=snap['items'][0]
            item.update(fullNameRu='Получатель 2',fullNameKz='Қабылдаушы 2',positionRu='Русская должность',positionKz='Қазақша лауазым',workplaceRu='Русская компания',workplaceKz='Қазақша ұйым')
            out=OUT/(tid+'-bilingual-regression.docx');render_docx(snap,out)
            with ZipFile(out) as z: text=''.join(E.fromstring(z.read('word/document.xml')).itertext())
            self.assertIn(item['fullNameRu'],text);self.assertIn(item['fullNameKz'],text)
            if tid!='biot-itr-certificate':
                self.assertIn(item['positionRu'],text);self.assertIn(item['positionKz'],text)
            if tid in ['biot-protocol','ptm-protocol','pb-protocol']:
                self.assertIn(item['workplaceRu'],text);self.assertIn(item['workplaceKz'],text)

    def test_15_numbered_cells_physical_pages_and_mutation(self):
        from print_contracts import assert_docx_columns,assert_pdf_columns,W
        from sanitize_templates import deterministic_zip
        from copy import deepcopy
        evidence=ROOT/'docs/evidence/commercial-acceptance/printing';evidence.mkdir(parents=True,exist_ok=True)
        snap=fixture('biot-protocol');snap['items'].append(deepcopy(snap['items'][0]));snap['items'][1].update(fullNameRu='Другой Получатель',protocolNumber='ПР-00002')
        good=evidence/'numbering-good.docx';render_docx(snap,good);convert_pdf(good,good.with_suffix('.pdf'))
        assert_docx_columns(good,'biot-protocol',2);checks=assert_pdf_columns(good.with_suffix('.pdf'),2,2)
        with ZipFile(good) as archive:files={n:archive.read(n) for n in archive.namelist()}
        tree=E.fromstring(files['word/document.xml']);tables=[t for t in tree.iter(W+'tbl') if len(t.findall(W+'tblGrid/'+W+'gridCol'))==7]
        for index,cell in enumerate(tables[1].findall(W+'tr')[1].findall(W+'tc'),8):next(cell.iter(W+'t')).text=str(index)
        files['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='utf-8');bad=evidence/'numbering-mutation-8-14.docx';deterministic_zip(bad,files);convert_pdf(bad,bad.with_suffix('.pdf'))
        with self.assertRaisesRegex(AssertionError,'DOCX_COLUMN_VALUES'):assert_docx_columns(bad,'biot-protocol',2)
        with self.assertRaisesRegex(AssertionError,'PDF_COLUMN_VALUES'):assert_pdf_columns(bad.with_suffix('.pdf'),2,2)
        (evidence/'numbering-mutation-result.json').write_text(json.dumps({'status':'PASS','method':'controlled mutation of second document column cells to 8..14, not historical-bug reproduction','good':checks,'mutantDocxRejected':True,'mutantPdfRejected':True},indent=2),encoding='utf8')

    def test_16_mapping_signature_basis_and_independent_dates(self):
        from print_contracts import W,text
        for tid in ['biot-protocol','ptm-protocol','pb-protocol','ps-protocol','ps-witness','biot-worker-card']:
            snap=fixture(tid);item=snap['items'][0];item['credentialNumber']='КР-98765';item['assignment']['education']='ТЕСТ: высшее'
            output=OUT/(tid+'-mapping-regression.docx');render_docx(snap,output)
            with ZipFile(output) as archive:tree=E.fromstring(archive.read('word/document.xml'))
            if tid in ['biot-protocol','ptm-protocol']:
                self.assertEqual(text(tree).count(snap['issuer']['approvalBasis']),1)
                table=next(tree.iter(W+'tbl'));cells=table.findall(W+'tr')[2 if tid=='biot-protocol' else 1].findall(W+'tc');self.assertEqual(text(cells[-1]).strip(),'')
            if tid=='ps-protocol':
                table=next(tree.iter(W+'tbl'));cell=table.findall(W+'tr')[1].findall(W+'tc')[-1];self.assertEqual(text(cell),'КР-98765')
            if tid=='pb-protocol':
                table=next(tree.iter(W+'tbl'));cell=table.findall(W+'tr')[1].findall(W+'tc')[3];self.assertEqual(text(cell),'ТЕСТ: высшее')
                import pdfplumber
                convert_pdf(output,output.with_suffix('.pdf'))
                with pdfplumber.open(output.with_suffix('.pdf')) as pdf:
                    tables=[t for t in pdf.pages[0].find_tables() if len(t.extract()[0])==5 and any('Образование' in (c or '') for c in t.extract()[0])]
                    self.assertEqual(len(tables),1);self.assertEqual(re.sub(r'\s+',' ',tables[0].extract()[1][3]).strip(),'ТЕСТ: высшее')
            if tid=='ps-witness':
                decision=next(p for p in tree.iter(W+'p') if 'Решением квалификационной' in text(p));self.assertIn('20',text(decision));self.assertNotIn('«22»',text(decision))

    def test_17_csv_semicolon_quoted_delimiters(self):
        path=OUT/'semicolon.csv';path.write_text('fullNameRu;fullNameKz;registrationNumber\n"Тестов; Иван";Әділбек;00012\n',encoding='utf8')
        result=import_table({'inputPath':str(path)},OUT/'semicolon-import.json')
        self.assertTrue(result['canApply']);self.assertEqual(result['rows'][0]['values']['fullNameRu'],'Тестов; Иван');self.assertEqual(result['rows'][0]['values']['registrationNumber'],'00012')

    def test_18_preflight_rejects_unbreakable_before_issuing(self):
        from renderer import preflight
        snapshots=[]
        for template in MANIFEST['templates']:
            snap=fixture(template['id']);snap['items'][0]['fullNameRu']='А'*81;snapshots.append(snap)
        result=preflight({'snapshots':snapshots},OUT/'preflight-overflow.json')
        self.assertEqual(result['issues'],[{'index':i,'code':'PRINT_LAYOUT_OVERFLOW'} for i in range(len(MANIFEST['templates']))])

    def test_19_numbered_header_repeated_on_physical_continuation(self):
        from print_contracts import assert_pdf_columns,W
        from sanitize_templates import deterministic_zip
        from copy import deepcopy
        evidence=ROOT/'docs/evidence/commercial-acceptance/printing';evidence.mkdir(parents=True,exist_ok=True)
        output=evidence/'numbering-physical-continuation.docx';render_docx(fixture('biot-protocol'),output)
        with ZipFile(output) as archive:files={n:archive.read(n) for n in archive.namelist()}
        tree=E.fromstring(files['word/document.xml']);table=next(tree.iter(W+'tbl'));source=table.findall(W+'tr')[2]
        for index in range(12):table.append(deepcopy(source))
        files['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='utf-8');deterministic_zip(output,files);convert_pdf(output,output.with_suffix('.pdf'))
        checks=assert_pdf_columns(output.with_suffix('.pdf'));self.assertGreater(len(checks),1)
        (evidence/'physical-continuation-result.json').write_text(json.dumps({'status':'PASS','method':'engineering-only duplicated table row forces table continuation; not a supported group protocol','checks':checks},indent=2),encoding='utf8')

    def test_20_formula_only_source_row_is_accounted(self):
        workbook=Workbook();worksheet=workbook.active;worksheet.append(['fullNameRu']);worksheet.append(['=1+1']);workbook.save(OUT/'formula-only.xlsx')
        result=import_table({'inputPath':str(OUT/'formula-only.xlsx')},OUT/'formula-only.json')
        self.assertFalse(result['canApply']);self.assertEqual(result['count'],1);self.assertEqual(result['rawRows'][0]['rowNumber'],2);self.assertIn('FORMULA_NOT_ALLOWED',result['rawRows'][0]['errors'])

    def test_24_biot_current_columns_numbers_results_and_preflight(self):
        from renderer import fields_for,preflight
        from print_contracts import W,text,assert_docx_columns
        from copy import deepcopy
        for tid,columns in [('biot-protocol',7),('biot-itr-protocol',10)]:
            snap=fixture(tid);item=snap['items'][0]
            item.update(number='BIOT-PROTOCOL-888888888888',protocolNumber='BIOT-PROTOCOL-888888888888',credentialNumber='BIOT-CERTIFICATE-888888888888')
            item['assignment'].update(biotKnowledgeResult='92 из 100',biotProctoringResult='прошел',result='сдал')
            output=OUT/(tid+'-current-contract.docx');render_docx(snap,output);assert_docx_columns(output,tid)
            with ZipFile(output) as archive:tree=E.fromstring(archive.read('word/document.xml'))
            table=next(tree.iter(W+'tbl'));cells=table.findall(W+'tr')[2].findall(W+'tc');self.assertEqual(len(cells),columns)
            fields=fields_for(snap,item)
            expected=['1',fields['WORKPLACE_BOTH'],fields['FULL_NAME_BOTH'],fields['POSITION_BOTH'],fields['DEPARTMENT_BOTH'],'сдал',''] if columns==7 else ['1',fields['EMPLOYER_BIN'],fields['EMPLOYER_NAME_ADDRESS_BOTH'],fields['FULL_NAME_BOTH'],fields['POSITION_BOTH'],'BIOT-CERTIFICATE-888888888888','92 из 100','прошел','сдал','']
            self.assertEqual([text(c).strip() for c in cells],expected)
            item['assignment']['biotUniqueNumber']='EXTERNAL-CERT-003';self.assertEqual(fields_for(snap,item)['BIOT_UNIQUE_NUMBER'],'EXTERNAL-CERT-003')
            item['assignment']['biotUniqueNumber']='';item['credentialNumber']='';self.assertEqual(fields_for(snap,item)['BIOT_UNIQUE_NUMBER'],'')
            item['assignment']['biotProctoringResult']='';self.assertEqual(fields_for(snap,item)['BIOT_PROCTORING_RESULT'],'')
        valid=fixture('biot-itr-protocol');valid['items'][0]['credentialNumber']='BIOT-CERTIFICATE-888888888888'
        too_wide=deepcopy(valid);too_wide['items'][0]['credentialNumber']='W'*70
        too_tall=fixture('biot-worker-card')
        for member in too_tall['issuer']['commission']:member.update(name='Представитель комиссии '*20,position='Должность члена комиссии '*20)
        result=preflight({'snapshots':[valid,too_wide,too_tall]},OUT/'biot-current-preflight.json')
        self.assertEqual(result['issues'],[{'index':1,'code':'PRINT_LAYOUT_OVERFLOW'},{'index':2,'code':'PRINT_LAYOUT_OVERFLOW'}])

    def test_21_word_package_causes_are_rejected(self):
        from print_contracts import assert_package_contract
        from sanitize_templates import deterministic_zip
        output=OUT/'word-contract.docx';render_docx(fixture('biot-protocol'),output);assert_package_contract(output)
        with ZipFile(output) as archive:files={n:archive.read(n) for n in archive.namelist()}
        tree=E.fromstring(files['word/document.xml']);tree.set('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable','undeclared')
        files['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='utf-8');mutant=OUT/'word-contract-mutant.docx';deterministic_zip(mutant,files)
        with self.assertRaisesRegex(AssertionError,'OOXML_UNDECLARED_MC_PREFIX'):assert_package_contract(mutant)

    def test_22_ps_title_not_painted_beneath_issuer_band(self):
        from print_contracts import assert_card_title_visible,W,text
        from sanitize_templates import deterministic_zip
        from upgrade_templates_v10 import move_box
        evidence=ROOT/'docs/evidence/commercial-acceptance/printing'
        good=evidence/'ps-title-visible.docx';snap=fixture('ps-card');render_docx(snap,good);convert_pdf(good,good.with_suffix('.pdf'))
        checks=assert_card_title_visible(good.with_suffix('.pdf'),'ps-card',1)
        with ZipFile(good) as archive:files={n:archive.read(n) for n in archive.namelist()}
        tree=E.fromstring(files['word/document.xml']);changed=0
        for box in tree.iter(W+'txbxContent'):
            if snap['items'][0]['fullNameRu'] in text(box):move_box(box,-9);changed+=1
        self.assertEqual(changed,2)
        files['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='utf-8');bad=evidence/'ps-title-hidden-mutation.docx';deterministic_zip(bad,files);convert_pdf(bad,bad.with_suffix('.pdf'))
        with self.assertRaisesRegex(AssertionError,'PS_TITLE_BEHIND_ISSUER_BAND'):assert_card_title_visible(bad.with_suffix('.pdf'),'ps-card',1)
        (evidence/'ps-title-visibility-result.json').write_text(json.dumps({'status':'PASS','good':checks,'mutantRejected':True},indent=2)+'\n',encoding='utf8')

    def test_23_biot_repeated_panel_and_issuer_divider(self):
        from print_contracts import assert_biot_card_panels,W,text
        from sanitize_templates import deterministic_zip
        from upgrade_templates_v10 import V,value,put
        from renderer import fields_for
        WP='{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
        A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
        evidence=ROOT/'docs/evidence/commercial-acceptance/printing'
        snap=fixture('biot-worker-card')
        # Keep this historic geometry regression pinned after the normative v16
        # replaces the repeating panels with the complete Appendix 4 form.
        import shutil
        historical=ROOT/'tests/render/fixtures/biot-worker-card.v15.docx'
        pinned=STORE/'biot-worker-card.v15.docx';shutil.copyfile(historical,pinned)
        snap.update(templateVersion=15,templateStorageKey=pinned.name,templateChecksum=hashlib.sha256(pinned.read_bytes()).hexdigest())
        snap['issuer']['commission']=[{'name':'Синтетический Председатель Әли','position':'Председатель контрольной комиссии'}]
        snap['items'][0]['assignment']['trainingSubject']='Тестовая программа безопасности и охраны труда для инженерного персонала'
        expected=fields_for(snap,snap['items'][0]);expected=[{k:expected[k] for k in ['CHAIR','SUBJECT']}]
        good=evidence/'biot-panel-good.docx';render_docx(snap,good);convert_pdf(good,good.with_suffix('.pdf'))
        checks=assert_biot_card_panels(good.with_suffix('.pdf'),1,expected)
        with ZipFile(good) as archive:original={n:archive.read(n) for n in archive.namelist()}
        tree=E.fromstring(original['word/document.xml'])
        boxes=[b for b in tree.iter(W+'txbxContent') if 'Синтетический Председатель' in text(b)]
        self.assertEqual(len(boxes),4)
        for box in boxes[-2:]:
            shape=next(n for n in box.iterancestors() if E.QName(n).localname in ['shape','anchor'])
            if shape.tag==V+'shape':shape.set('style',put(put(shape.get('style'),'margin-left',282.9),'width',259.5))
            else:
                shape.find(WP+'positionH/'+WP+'posOffset').text='3592830'
                shape.find(WP+'extent').set('cx','3295650');shape.find('.//'+A+'xfrm/'+A+'ext').set('cx','3295650')
        files=dict(original);files['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='utf-8')
        bad=evidence/'biot-panel-wide-mutation.docx';deterministic_zip(bad,files);convert_pdf(bad,bad.with_suffix('.pdf'))
        with self.assertRaisesRegex(AssertionError,'BIOT_RIGHT_PANEL_(CLIP|INSET)'):assert_biot_card_panels(bad.with_suffix('.pdf'),1,expected)
        tree=E.fromstring(original['word/document.xml']);changed=0
        for anchor in tree.iter(WP+'anchor'):
            extent=anchor.find(WP+'extent')
            if int(extent.get('cx','0'))!=4445:continue
            offset=anchor.find(WP+'positionV/'+WP+'posOffset');offset.text=str(int(offset.text)-14*12700)
            extent.set('cy',str(int(extent.get('cy'))+14*12700));extent=anchor.find('.//'+A+'xfrm/'+A+'ext');extent.set('cy',str(int(extent.get('cy'))+14*12700));changed+=1
        for shape in tree.iter(V+'shape'):
            style=shape.get('style','')
            if abs(value(style,'width')-.35)>.001 or value(style,'height')<150:continue
            shape.set('style',put(put(style,'margin-top',value(style,'margin-top')-14),'height',value(style,'height')+14));changed+=1
        self.assertEqual(changed,4)
        files=dict(original);files['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='utf-8')
        crossing=evidence/'biot-divider-crossing-mutation.docx';deterministic_zip(crossing,files);convert_pdf(crossing,crossing.with_suffix('.pdf'))
        with self.assertRaisesRegex(AssertionError,'BIOT_DIVIDER_THROUGH_ISSUER'):assert_biot_card_panels(crossing.with_suffix('.pdf'),1,expected)
        (evidence/'biot-panel-regression-result.json').write_text(json.dumps({'status':'PASS','good':checks,'widePanelMutantRejected':True,'dividerMutantRejected':True,'longChairAndProgram':expected},ensure_ascii=False,indent=2)+'\n',encoding='utf8')

    def test_12_photo_exif_and_decompression_limits(self):
        image=Image.new('RGB',(480,640),'#aabecc');exif=Image.Exif();exif[274]=6;image.save(OUT/'rotated.jpg',exif=exif)
        result=photo_normalize(OUT/'rotated.jpg',OUT/'rotated-normalized.png',{});self.assertEqual((result['width'],result['height']),(640,480))
        Image.new('RGB',(5000,5000),'white').save(OUT/'pixel-limit.png')
        with self.assertRaisesRegex(ValueError,'PHOTO_PIXEL_LIMIT'):photo_normalize(OUT/'pixel-limit.png',OUT/'invalid-photo.png',{})
        (OUT/'byte-limit.png').write_bytes(b'0'*(5*1024*1024+1))
        with self.assertRaisesRegex(ValueError,'PHOTO_BYTES_LIMIT'):photo_normalize(OUT/'byte-limit.png',OUT/'invalid-photo.png',{})

if __name__=='__main__':unittest.main(verbosity=2)
