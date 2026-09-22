import io
import sys
import time
import unittest
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.parsers import expat

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'scripts/render'))
from xml_input import validate_xlsx_xml


def archive(xml, name='xl/worksheets/sheet1.xml'):
    output = io.BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as z:
        z.writestr(name, xml)
    output.seek(0)
    return ZipFile(output)


class XmlInputTests(unittest.TestCase):
    def test_upstream_dtd_attribute_case_is_rejected_before_normalization(self):
        # Same precondition as upstream PR1321, with a deliberately small baseline.
        n = 1000
        definitions = ''.join(f'<!ATTLIST e a{i} NMTOKEN "x">' for i in range(n))
        attributes = ' '.join(f'a{i}=" v "' for i in range(n))
        xml = f'<!DOCTYPE e [{definitions}]><e {attributes}/>'.encode()
        baseline = expat.ParserCreate()
        seen = []
        baseline.StartElementHandler = lambda _name, attrs: seen.append(attrs['a0'])
        baseline.Parse(xml, True)
        self.assertEqual(seen, ['v'])  # Baseline reaches vulnerable normalization.
        with archive(xml) as z, self.assertRaisesRegex(ValueError, 'DTD_FORBIDDEN'):
            validate_xlsx_xml(z)

    def test_large_attribute_attack_rejected_at_doctype(self):
        n = 100_000
        xml = ('<!DOCTYPE e [' + ''.join(f'<!ATTLIST e a{i} NMTOKEN "x">' for i in range(n))
               + ']><e ' + ' '.join(f'a{i}=" v "' for i in range(n)) + '/>').encode()
        started = time.monotonic()
        with archive(xml) as z, self.assertRaisesRegex(ValueError, 'DTD_FORBIDDEN'):
            validate_xlsx_xml(z)
        self.assertLess(time.monotonic() - started, 5)

    def test_utf16_and_chunk_boundary_doctypes(self):
        for xml in ['<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE e [<!ATTLIST e a NMTOKEN "x">]><e a=" v "/>'.encode('utf-16'),
                    (b' ' * (64 * 1024 - 4)) + b'<!DOCTYPE e [<!ATTLIST e a NMTOKEN "x">]><e a=" v "/>']:
            with self.subTest(size=len(xml)), archive(xml) as z, self.assertRaisesRegex(ValueError, 'DTD_FORBIDDEN'):
                validate_xlsx_xml(z)

    def test_all_xml_parts_and_relationships_are_checked(self):
        for name in ['[Content_Types].xml', 'docProps/core.xml', 'xl/_rels/workbook.xml.rels']:
            with self.subTest(part=name), archive(b'<!DOCTYPE e SYSTEM "file:///synthetic-never-read"><e/>', name) as z:
                with self.assertRaisesRegex(ValueError, 'DTD_FORBIDDEN'):
                    validate_xlsx_xml(z)

    def test_depth_and_attribute_bounds(self):
        for xml, message in [(b'<e>' * 129 + b'</e>' * 129, 'DEPTH_LIMIT'),
                             (('<e ' + ' '.join(f'a{i}="v"' for i in range(129)) + '/>').encode(), 'ATTRIBUTE_LIMIT')]:
            with archive(xml) as z, self.assertRaisesRegex(ValueError, message):
                validate_xlsx_xml(z)

    def test_valid_openpyxl_workbook_roundtrip(self):
        from openpyxl import Workbook, load_workbook
        workbook = Workbook()
        workbook.active.append(['ФИО', 'Организация', 'Должность'])
        workbook.active.append(['Ә Ғ Қ Ң Ө Ұ Ү Һ І', 'Учебный центр', 'Инженер'])
        output = io.BytesIO()
        workbook.save(output)
        with ZipFile(output) as z:
            validate_xlsx_xml(z)
        output.seek(0)
        restored = load_workbook(output, read_only=True, data_only=False, keep_links=False)
        self.assertEqual(list(restored.active.values)[1][0], 'Ә Ғ Қ Ң Ө Ұ Ү Һ І')
        restored.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
