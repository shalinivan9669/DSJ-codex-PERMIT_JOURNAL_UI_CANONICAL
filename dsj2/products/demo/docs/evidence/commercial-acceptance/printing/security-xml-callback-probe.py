"""Bounded diagnostic of the actual openpyxl/lxml cell serializer callback sizes."""
from pathlib import Path
import hashlib,inspect,json,sys
import lxml.etree as E
import openpyxl
from openpyxl.worksheet._writer import WorksheetWriter
import openpyxl.cell._writer as cell_writer
ROOT=Path.cwd();OUT=ROOT/'docs/evidence/commercial-acceptance/printing'
class Sink:
 def __init__(self):self.total=0;self.maximum=0;self.calls=0
 def write(self,b):self.total+=len(b);self.maximum=max(self.maximum,len(b));self.calls+=1;return len(b)
sink=Sink();wb=openpyxl.Workbook();ws=wb.active
for i in range(1000):
 ws.append(['<&Ә'*166+'Қң']*28)
 for c in ws[ws.max_row]:c.data_type='s'
writer=WorksheetWriter(ws,out=sink);writer.write()
assert sink.total>10_000_000 and sink.maximum<1_000_000
paths=[Path(inspect.getfile(WorksheetWriter)),Path(inspect.getfile(cell_writer))]
report={'status':'PASS','python':sys.version,'lxml':E.LXML_VERSION,'libxml':E.LIBXML_VERSION,'openpyxl':openpyxl.__version__,'rows':1000,'columns':28,'cellCharacters':500,'totalXmlBytes':sink.total,'maximumCallbackBytes':sink.maximum,'callbackCount':sink.calls,'limitations':'Bounded measurement complements actual writer source review; not an exploit test, native patch claim, or proof for arbitrary application callbacks. Total worksheet size is distinct from a single buffered output callback.','sourceHashes':{str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},'sourceReviewed':{'get_stream':inspect.getsource(WorksheetWriter.get_stream),'write_rows':inspect.getsource(WorksheetWriter.write_rows),'write_row':inspect.getsource(WorksheetWriter.write_row),'lxml_write_cell':inspect.getsource(cell_writer.lxml_write_cell)}}
(OUT/'security-xml-callback-probe.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps({k:v for k,v in report.items() if k not in ['sourceHashes','sourceReviewed']},ensure_ascii=False))
